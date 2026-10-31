"""模板答卷的构建、内容基线与只读候选提取。"""

import json
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


META_SHEET = '_dy_qa'
LINKS_SHEET = '_dy_links'
MATERIALS_SHEET = '已确认材料'
MATERIAL_HEADERS = ('配置项', '调研问题', '材料名称', '已确认要求')


def answer_column(template):
    return template['layout']['headers'].index('答案') + 1


def _text_cell(sheet, row, column, value):
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        cell.data_type = 's'
    cell.alignment = Alignment(vertical='top', wrap_text=True)
    return cell


def _make_book(template, records, metadata):
    book = Workbook()
    sheet = book.active
    sheet.title = '调研清单'
    column = answer_column(template)
    for row, values in enumerate(template['layout']['rows'], 1):
        for col, value in enumerate(values[:column - 1], 1):
            _text_cell(sheet, row, col, value)
    for merge in template['layout']['merges']:
        sheet.merge_cells(merge)
    for letter, width in template['layout']['widths'].items():
        sheet.column_dimensions[letter].width = width
    for offset, header in enumerate(('已确认答案', '本轮补充或修订', '调研状态')):
        cell = _text_cell(sheet, template['layout']['header_row'], column + offset, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D9EAF7')
        sheet.column_dimensions[get_column_letter(column + offset)].width = 42 if offset < 2 else 16
    by_id = {record['question_id']: record for record in records}
    links = []
    _text_cell(sheet, template['layout']['header_row'], 1, '序号')
    for number, question in enumerate(template['questions'], 1):
        _text_cell(sheet, question['row'], 1, number)
        rows = question.get('rows', [question['row']])
        for col in range(column, column + 3):
            if len(rows) > 1:
                sheet.merge_cells(start_row=rows[0], start_column=col, end_row=rows[-1], end_column=col)
        record = by_id.get(question['question_id'])
        _text_cell(sheet, question['row'], column, record['answer'] if record else None)
        _text_cell(sheet, question['row'], column + 1, None)
        status = ('已完成' if record['research_complete'] else '部分完成') if record else '未回答'
        _text_cell(sheet, question['row'], column + 2, status)
    sheet.freeze_panes = f'A{template["layout"]["header_row"] + 1}'
    for attachment in template['attachments']:
        detail = book.create_sheet(attachment['name'])
        for col, header in enumerate(attachment['headers'], 1):
            _text_cell(detail, 1, col, header)
            detail.column_dimensions[get_column_letter(col)].width = 24
        # 字段明细按原列投影；每行保留所属问题及材料，多个绑定顺序追加。
        linked = []
        for question in template['questions']:
            if attachment['name'] not in question['attachment_names']:
                continue
            for bound in by_id.get(question['question_id'], {}).get('attachments', []):
                rows = [row for page in bound.get('details', []) if page['sheet'] == attachment['name'] for row in page['rows']]
                if rows:
                    linked.extend((question, bound, row, '\n'.join(bound['requirements'])) for row in rows)
                else:
                    linked.extend((question, bound, {}, requirement) for requirement in bound['requirements'])
        if linked:
            for offset, header in enumerate(MATERIAL_HEADERS, len(attachment['headers']) + 1):
                _text_cell(detail, 1, offset, header)
            for detail_row, (question, bound, values, requirement) in enumerate(linked, 2):
                for col, header in enumerate(attachment['headers'], 1):
                    _text_cell(detail, detail_row, col, None if values.get(header) == '' else values.get(header))
                for offset, value in enumerate((question['config_item'], question['question'], Path(bound['reference']).name, requirement), len(attachment['headers']) + 1):
                    _text_cell(detail, detail_row, offset, value)
                links.append((attachment['name'], detail_row, question['question_id'], bound['reference'], bound['sha256']))
    # 通用页覆盖全部确认绑定，包括没有模板附件页的问题。
    materials = book.create_sheet(MATERIALS_SHEET)
    for col, header in enumerate(MATERIAL_HEADERS, 1):
        _text_cell(materials, 1, col, header)
        materials.column_dimensions[get_column_letter(col)].width = 40
    material_row = 2
    for record in records:
        question = next(q for q in template['questions'] if q['question_id'] == record['question_id'])
        for bound in record['attachments']:
            for requirement in bound['requirements']:
                for col, value in enumerate((question['config_item'], question['question'], Path(bound['reference']).name, requirement), 1):
                    _text_cell(materials, material_row, col, value)
                links.append((MATERIALS_SHEET, material_row, record['question_id'], bound['reference'], bound['sha256']))
                material_row += 1
    mapping = book.create_sheet(LINKS_SHEET)
    for row, values in enumerate([('sheet', 'row', 'question_id', 'reference', 'sha256'), *links], 1):
        for col, value in enumerate(values, 1):
            _text_cell(mapping, row, col, value)
    mapping.sheet_state = 'veryHidden'
    meta = book.create_sheet(META_SHEET)
    _text_cell(meta, 1, 1, json.dumps(metadata, ensure_ascii=False, sort_keys=True))
    meta.sheet_state = 'veryHidden'
    return book


def build_workbook(path, template, records, metadata):
    """写入一个新工作簿并返回语义基线，不使用不可编辑的整文件摘要。"""
    book = _make_book(template, records, metadata)
    book.save(path)
    return describe(_load(path))


def _json_scalar(value):
    if isinstance(value, timedelta):
        # Excel 累计时长保留为 ISO 8601 duration，避免超过24小时被误当时刻。
        sign = '-' if value < timedelta(0) else ''
        span = abs(value)
        hours, seconds = divmod(span.seconds, 3600)
        minutes, seconds = divmod(seconds, 60)
        fraction = f'.{span.microseconds:06d}'.rstrip('0') if span.microseconds else ''
        return f'{sign}P{span.days}DT{hours}H{minutes}M{seconds}{fraction}S'
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def describe(book):
    return [{'name': sheet.title,
             'merges': sorted(str(merge) for merge in sheet.merged_cells.ranges),
             'rows': [[_json_scalar(cell.value) for cell in row] for row in sheet.iter_rows()]}
            for sheet in book.worksheets]


def _load(path):
    try:
        return load_workbook(path, data_only=False)
    except Exception as exc:
        raise ValueError('答卷无法读取或格式损坏') from exc


def verify_snapshot(path, template, records, metadata):
    """不可变答卷必须与同版机器记忆完全一致。"""
    actual = describe(_load(path))
    if actual != describe(_make_book(template, records, metadata)):
        raise ValueError('同版本答卷与记忆内容不一致')
    return actual


def metadata_from(path):
    book = _load(path)
    try:
        metadata = json.loads(book[META_SHEET]['A1'].value)
    except (KeyError, ValueError, TypeError) as exc:
        raise ValueError('拒绝旧格式或缺少登记元数据的答卷') from exc
    if not isinstance(metadata, dict):
        raise ValueError('答卷元数据格式无效')
    return book, metadata


def _value(rows, row, column):
    return rows[row - 1][column - 1] if row <= len(rows) and column <= len(rows[row - 1]) else None


def extract_candidates(book, template, registration):
    """只提取允许编辑的本轮栏和附件候选，不确认、不写磁盘。"""
    baseline = registration['baseline']
    actual = describe(book)
    if [sheet['name'] for sheet in actual] != [sheet['name'] for sheet in baseline]:
        raise ValueError('答卷工作表被修改')
    for sheet in book:
        for row in sheet:
            if any(cell.data_type == 'f' for cell in row):
                raise ValueError('答卷不接受公式，请填写明确文本')
    column = answer_column(template)
    question_rows = {question['row']: question for question in template['questions']}
    question_ids = {question['question_id'] for question in template['questions']}
    original_links = next(page for page in baseline if page['name'] == LINKS_SHEET)
    current_links = next(page for page in actual if page['name'] == LINKS_SHEET)
    if original_links != current_links:
        raise ValueError('答卷隐藏关联被修改')
    bindings = {(row[0], row[1]): row[2] for row in original_links['rows'][1:]}
    answers, changes = [], []
    for index, (current, original) in enumerate(zip(actual, baseline, strict=True)):
        if current['merges'] != original['merges']:
            raise ValueError('答卷合并结构被修改')
        max_row = max(len(current['rows']), len(original['rows']))
        max_col = max(max(map(len, current['rows']), default=0), max(map(len, original['rows']), default=0))
        for row in range(1, max_row + 1):
            for col in range(1, max_col + 1):
                value = _value(current['rows'], row, col)
                previous = _value(original['rows'], row, col)
                if value == previous:
                    continue
                if index == 0 and col == column + 1 and row in question_rows:
                    if value is not None and str(value).strip():
                        question_id = question_rows[row]['question_id']
                        answers.append({'question_id': question_id, 'answer': str(value).strip(),
                                        'research_complete': registration['completeness'].get(question_id, False),
                                        'expected_revision': registration['revisions'].get(question_id)})
                elif current['name'] not in (actual[0]['name'], META_SHEET, LINKS_SHEET) and row > 1:
                    linked_qid = bindings.get((current['name'], row))
                    changes.append({'sheet': current['name'], 'cell': f'{get_column_letter(col)}{row}',
                                    'row': row, 'column': col, 'before': previous, 'value': value,
                                    'scope_id': registration['metadata']['scope_id'],
                                    'question_ids': ([linked_qid] if linked_qid in question_ids else
                                                     [q['question_id'] for q in template['questions']
                                                      if current['name'] in q['attachment_names']]),
                                    'can_write_confirmed': False})
                else:
                    raise ValueError('答卷题目、已确认栏、状态、附件表头或元数据被修改')
    return answers, changes
