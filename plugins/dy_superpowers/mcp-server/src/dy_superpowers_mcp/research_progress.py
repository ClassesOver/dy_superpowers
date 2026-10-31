"""旧研究存储的历史覆盖率与旧答卷实现；新版业务入口直接使用 qa_store。"""

from __future__ import annotations

import re
import json
from io import BytesIO
from datetime import datetime
from pathlib import Path

from .memory import validate_module
from .research import record_summary
from .storage import record_lock


HEADERS = ["模块", "检查顺序", "配置项", "调研子项", "配置子项", "调研问题",
           "回答指引与提取要素", "已记忆项", "当前状态", "答案", "附件/补充材料"]
MODULE_LABELS = {"general-ledger": "总账", "expense": "费用"}


def _guide_items(module: str) -> list[dict[str, str]]:
    if module not in {"expense", "general-ledger"}:
        raise ValueError("当前只支持费用和总账调研")
    guide = Path(__file__).parents[3] / "skills/module-research/references" / f"{module}-research-guide.md"
    questions, items = {}, {}
    qa_header = matrix_header = None
    for line in guide.read_text(encoding="utf-8").splitlines():
        if line.startswith("#"):
            qa_header = matrix_header = None
        if not line.startswith("| ") or not line.endswith("|"):
            continue
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        if "配置项" in cells and any(label in cells for label in ("汇总调研问题", "汇总问题")):
            qa_header = cells
            matrix_header = None
            continue
        if "补充汇总问题" in cells and "检查顺序" in cells:
            qa_header = cells
            matrix_header = None
            continue
        if "模块" in cells and "子项" in cells and "提取指引要素" in cells:
            matrix_header = cells
            qa_header = None
            continue
        if qa_header and len(cells) >= len(qa_header):
            order_col = qa_header.index("检查编号") if "检查编号" in qa_header else qa_header.index("检查顺序")
            order = cells[order_col]
            questions[order] = {
                "config_item": cells[qa_header.index("配置项")] if "配置项" in qa_header else "",
                "question": cells[next(qa_header.index(label) for label in ("汇总调研问题", "汇总问题", "补充汇总问题") if label in qa_header)],
            }
            continue
        if not matrix_header or len(cells) < len(matrix_header) or cells[0] != module:
            continue
        order, research_id = cells[matrix_header.index("检查顺序")], cells[matrix_header.index("子项")]
        extraction = cells[matrix_header.index("提取指引要素")]
        attachment = cells[matrix_header.index("关联附件清单") if "关联附件清单" in matrix_header else matrix_header.index("关联附件模板")]
        internal_order = order
        if module == "expense":
            match = re.fullmatch(r"EXP-(\d{3})", order)
            if not match:
                continue
            internal_order = str(int(match.group(1)))
        group = questions.get(order, {})
        config_subitem = re.split(r"[：:]", extraction, maxsplit=1)[0].strip()
        if not group.get("config_item"):
            group["config_item"] = config_subitem
        items[internal_order, research_id] = {
            "check_order": internal_order, "display_order": order, "research_id": research_id,
            "question": group.get("question", ""), "attachment": attachment.replace("<br>", "、"),
            "config_item": group.get("config_item", ""),
            "config_subitem": config_subitem,
            "extraction_guide": extraction,
        }
    if not items:
        raise ValueError(f"{module} 调研指南缺少矩阵映射子项")
    return list(items.values())


def calculate_progress(module: str, records: list[dict], *, scope: str | None = None,
                       scope_id: str | None = None) -> dict:
    module = validate_module(module)
    if module not in {"expense", "general-ledger"}:
        raise ValueError("当前只支持费用和总账调研")
    for name, value in (("scope", scope), ("scope_id", scope_id)):
        if value is not None and (not isinstance(value, str) or not value.strip()):
            raise ValueError(f"{name} 必须是非空文本")
    module_records = [record for record in records if record.get("module") == module]
    if scope is None and scope_id is None:
        scopes = {record.get("scope_id") for record in module_records if record.get("scope_id")}
        if len(scopes) > 1:
            raise ValueError("该模块存在多个适用范围；请指定 scope 或 scope_id，不能跨范围合并")
    selected = [record for record in module_records
                if (scope is None or record.get("scope") == scope)
                and (scope_id is None or record.get("scope_id") == scope_id)]
    if scope_id is None and scope is not None:
        scope_ids = {record.get("scope_id") for record in selected if record.get("scope_id")}
        if len(scope_ids) > 1:
            raise ValueError("该适用范围关联多个稳定范围标识；请指定 scope_id，不能合并计算")
    items = _guide_items(module)
    expected = {(item["check_order"], item["research_id"]) for item in items}
    confirmed_records = [record for record in selected if record.get("status") == "confirmed"]
    confirmed = {(record.get("check_order"), record.get("research_id"))
                 for record in confirmed_records} & expected
    answered = {(record.get("check_order"), record.get("research_id")) for record in confirmed_records
                if record.get("research_complete") is True} & expected
    partial = ({(record.get("check_order"), record.get("research_id")) for record in confirmed_records
                if record.get("research_complete") is False} & expected) - answered
    unknown = confirmed - answered - partial
    percentage = len(answered) / len(items) * 100 if items else 0
    confirmed_percentage = len(confirmed) / len(items) * 100 if items else 0
    unanswered_count = len(items) - len(confirmed)
    missing = [item for item in items if (item["check_order"], item["research_id"]) not in answered]
    return {"module": module, "scope": scope, "scope_id": scope_id,
            "answered": len(answered), "total": len(items), "percentage": percentage,
            "confirmed_count": len(confirmed), "confirmed_percentage": confirmed_percentage,
            "partial_count": len(partial), "unknown_count": len(unknown),
            "unanswered_count": unanswered_count,
            "progress_summary": (
                f"已确认回答 {len(confirmed)}/{len(items)}（{confirmed_percentage:.1f}%）；"
                f"完整完成 {len(answered)}/{len(items)}（{percentage:.1f}%）；"
                f"部分完成 {len(partial)} 项；完整性待核实 {len(unknown)} 项；"
                f"尚未回答 {unanswered_count} 项。"),
            "route": "interactive" if percentage >= 70 else "excel", "items": missing,
            "unmapped_records": [record_summary(record) for record in selected
                                 if (record.get("check_order"), record.get("research_id")) not in expected],
            "completeness_unknown": [record_summary(record) for record in selected
                                     if record.get("status") == "confirmed"
                                     and "research_complete" not in record]}


def calculate_routing_progress(module: str, records: list[dict], *, scope: str | None = None,
                               scope_id: str | None = None) -> dict:
    """未选范围也可分流；多范围返回独立进度及不扣减答案的通用清单。"""
    if scope is not None or scope_id is not None:
        result = calculate_progress(module, records, scope=scope, scope_id=scope_id)
        if scope is None:
            result["scope"] = next((record.get("scope") for record in records
                                    if record.get("module") == module and record.get("scope_id") == scope_id), None)
        result["scope_mode"] = "selected"
        return result

    result = calculate_progress(module, [])
    scope_ids = sorted({record["scope_id"] for record in records if record.get("module") == module})
    if len(scope_ids) == 1:
        result = calculate_progress(module, records, scope_id=scope_ids[0])
        result["scope"] = next(record.get("scope") for record in records
                               if record.get("module") == module and record["scope_id"] == scope_ids[0])
        result["scope_mode"] = "selected"
    elif scope_ids:
        result["scope_mode"] = "multiple"
        result["scope_progress"] = []
        for selected_id in scope_ids:
            selected = calculate_progress(module, records, scope_id=selected_id)
            result["scope_progress"].append({
                "scope_id": selected_id,
                "scope": next(record.get("scope") for record in records
                              if record.get("module") == module and record["scope_id"] == selected_id),
                **{key: selected[key] for key in (
                    "answered", "total", "percentage", "route", "confirmed_count", "confirmed_percentage",
                    "partial_count", "unknown_count", "unanswered_count", "progress_summary")},
            })
            for field in ("unmapped_records", "completeness_unknown"):
                result[field].extend(selected[field])
        result["progress_summary"] = "未选定范围；请分别展示 scope_progress 中各范围的已确认回答和完整完成进度，不合并计数。"
    else:
        result["scope_mode"] = "unscoped"
    return result


def export_unanswered_template(workspace: str | Path, module: str, trace_id: str, *,
                               records: list[dict], scope: str | None = None, scope_id: str | None = None,
                               confirmed: bool) -> dict:
    if confirmed is not True:
        raise PermissionError("导出待补问题清单需要用户明确确认")
    if not isinstance(trace_id, str) or not re.fullmatch(r"[A-Za-z0-9_-]{1,128}", trace_id):
        raise ValueError("trace_id 格式不合法")
    progress = calculate_routing_progress(module, records, scope=scope, scope_id=scope_id)
    module = progress["module"]
    output_dir = Path(workspace) / ".dy_superpowers" / "research" / module
    workspace_path = Path(workspace).resolve()
    current = Path(workspace)
    for part in (".dy_superpowers", "research", module):
        current = current / part
        if current.is_symlink():
            raise ValueError("待补问题清单目录不得包含符号链接")
    output_dir.mkdir(parents=True, exist_ok=True)
    if not output_dir.resolve().is_relative_to(workspace_path):
        raise ValueError("待补问题清单必须保存在当前工作区")
    items = _guide_items(module)
    selected = [record for record in records if record.get("module") == module
                and (scope is None or record.get("scope") == scope)
                and (scope_id is None or record.get("scope_id") == scope_id)]
    if progress["scope_mode"] == "multiple" and scope is None and scope_id is None:
        selected = [record for record in records if record.get("module") == module]
    memory = {}
    for record in selected:
        key = (str(record.get("check_order")), str(record.get("research_id")))
        if key in {(item["check_order"], item["research_id"]) for item in items}:
            memory.setdefault(key, []).append(record)
    archived_sheets = _latest_attachment_sheets(workspace, module)
    workbook = _new_research_workbook(module, items, memory, archived_sheets)
    payload = BytesIO()
    workbook.save(payload)
    workbook.close()
    label = MODULE_LABELS[module]
    date = datetime.now().astimezone().strftime("%Y%m%d")
    stem = f"{label}调研清单{date}"
    lock_target = output_dir / ".serial"
    if lock_target.is_symlink() or lock_target.with_suffix(".lock").is_symlink():
        raise ValueError("调研清单导出锁不得为符号链接")
    with record_lock(lock_target):
        existing = [int(match.group(1)) for path in output_dir.iterdir()
                    if not path.is_symlink()
                    and (match := re.fullmatch(re.escape(stem) + r"(\d{2,})\.xlsx", path.name))]
        serial = max(existing, default=0) + 1
        while True:
            output = output_dir / f"{stem}{serial:02d}.xlsx"
            try:
                with output.open("xb") as stream:
                    stream.write(payload.getvalue())
                break
            except FileExistsError:
                serial += 1
    return {"module": module, "path": str(output.relative_to(workspace)),
            "question_count": len(progress["items"]), "scope": progress["scope"],
            "scope_id": progress["scope_id"], "scope_mode": progress["scope_mode"]}


def _question_rows(module: str, items: list[dict], memory: dict) -> list[list]:
    rows = []
    for item in items:
        key = (item["check_order"], item["research_id"])
        records = [record for record in memory.get(key, []) if record.get("status") == "confirmed"]
        memos = [f"【{record.get('scope')}】{record['conclusion']}" if record.get("scope")
                 else str(record["conclusion"]) for record in records]
        if any(record.get("research_complete") is not True for record in records):
            status = "已有局部确认事实；完整性待补充"
        elif records:
            status = "已确认"
        else:
            status = "待调研"
        attachment = item["attachment"]
        rows.append([MODULE_LABELS[module], item["display_order"], item["config_item"],
                     item["research_id"], item["config_subitem"], item["question"],
                     item["extraction_guide"], "\n".join(memos) or None, status, None,
                     "附件：" + attachment if attachment and attachment != "无" else None])
    return rows


def _new_research_workbook(module: str, items: list[dict], memory: dict,
                           archived_sheets: dict | None):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles.colors import Color
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "调研问题清单"
    templates = _attachment_templates(module)
    for name in dict.fromkeys([*templates, *(archived_sheets or {})]):
        detail = workbook.create_sheet(name)
        saved = (archived_sheets or {}).get(name)
        detail.append(saved["headers"] if saved is not None else templates[name])
        if saved is not None:
            for values in saved["rows"]:
                detail.append(values)
        if saved is None or (saved.get("example_present") and not saved["rows"]):
            detail.append([f"示例{field}" for field in templates.get(name, saved["headers"] if saved else [])])
        _style_attachment_sheet(detail)

    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color=Color(theme=0), name="宋体", size=11)
        cell.fill = PatternFill("solid", fgColor="FF2F6B5F")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    rows = _question_rows(module, items, memory)
    for row_number, row in enumerate(rows, 2):
        sheet.append(row)
        key = (items[row_number - 2]["check_order"], items[row_number - 2]["research_id"])
        scopes = [{"scope": record.get("scope", ""), "scope_id": record.get("scope_id", ""),
                   "research_complete": record.get("research_complete") is True}
                  for record in memory.get(key, []) if record.get("status") == "confirmed"]
        if scopes:
            sheet.cell(row_number, 8).comment = Comment(
                json.dumps(scopes, ensure_ascii=False), "dy_superpowers")
    _merge_question_groups(sheet, rows)
    widths = [8.78, 14.89, 14.89, 11.11, 14.89, 14.89, 27.56, 18.78, 18.78, 16.67, 18.78]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2, max_col=len(HEADERS)):
        sheet.row_dimensions[row[0].row].height = 88
        for cell in row:
            cell.font = Font(name="宋体", size=11)
            vertical, horizontal = "top", None
            if cell.column in (1, 2, 3):
                vertical, horizontal = "center", "center"
            elif cell.column == 6:
                vertical, horizontal = "center", "left"
            elif cell.column == 7:
                vertical = "center"
            elif cell.column == 11:
                cell.font = Font(name="宋体", size=10, color="FF0563C1")
                cell.fill = PatternFill("solid", fgColor="FFFFF8E8")
            cell.alignment = Alignment(horizontal=horizontal, vertical=vertical, wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:K{sheet.max_row}"
    sheet.sheet_format.defaultRowHeight = 14.4

    return workbook


def _style_attachment_sheet(sheet) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF24476B")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _merge_question_groups(sheet, rows: list[list]) -> None:
    start = 2
    while start < len(rows) + 2:
        end = start
        while end < len(rows) + 1 and rows[end - 2][1] == rows[end - 1][1]:
            end += 1
        for column in (1, 2, 3, 6):
            if end > start:
                sheet.merge_cells(start_row=start, start_column=column, end_row=end, end_column=column)
        start = end + 1


def _attachment_templates(module: str) -> dict[str, list[str]]:
    guide = Path(__file__).parents[3] / "skills/module-research/references" / f"{module}-research-guide.md"
    templates = {}
    current = None
    for line in guide.read_text(encoding="utf-8").splitlines():
        if line.startswith("### 模板："):
            current = line.removeprefix("### 模板：").strip()
        elif line.startswith("#"):
            current = None
        if module == "general-ledger" and current and line.startswith("| "):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if len(cells) == 2 and cells[0] != "字段" and not re.fullmatch(r":?-{3,}:?", cells[0]):
                templates.setdefault(current, []).append(cells[0])
            continue
        if not line.startswith("| ") or not line.endswith("|"):
            continue
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if len(cells) < 2 or not re.fullmatch(r"\d{2}_.+", cells[0]):
            continue
        fields = [field.strip() for field in re.split(r"[；;]", cells[1]) if field.strip()]
        templates[cells[0]] = fields
    return templates


def _latest_attachment_sheets(workspace: str | Path, module: str) -> dict | None:
    from .source_archive import _validated_run

    root = Path(workspace).resolve()
    index_path = root / ".dy_superpowers-parsed/index.json"
    if index_path.is_symlink() or not index_path.is_file():
        return None
    index = json.loads(index_path.read_text(encoding="utf-8"))
    archived = []
    for digest, source in index.get("sources", {}).items():
        relative = source.get("latest_complete")
        if not relative:
            continue
        manifest = _validated_run(root, relative, digest)
        options = manifest.get("options", {})
        if options.get("module") != module or options.get("suffix") not in {".xlsx", ".xlsm"}:
            continue
        run = root / relative
        candidates_path = run / "candidates.json"
        if candidates_path.is_symlink() or not candidates_path.is_file():
            continue
        data = json.loads(candidates_path.read_text(encoding="utf-8"))
        if isinstance(data.get("attachment_sheets"), dict):
            archived.append((manifest["created_at"], data["attachment_sheets"]))
    return max(archived, default=("", None), key=lambda value: value[0])[1]
