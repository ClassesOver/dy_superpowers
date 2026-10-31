"""聊天修订与同模板旧答卷续答合并后，完整走 prepare/报告保存回读。"""
import hashlib
from pathlib import Path
import pytest
from openpyxl import load_workbook
import server
from dy_superpowers_mcp.qa_templates import load_template
from dy_superpowers_mcp.qa_workbook import answer_column
from template_qa_helpers import make_report


@pytest.mark.parametrize('module',['expense','general-ledger'])
def test_chat_old_workbook_merge_and_current_report_delivery(tmp_path,monkeypatch,module):
    confirmed_col=answer_column(load_template(module));new_col=confirmed_col+1
    monkeypatch.setenv('CODEX_HOME',str(tmp_path/'codex'))
    monkeypatch.delenv('DB_NAME',raising=False)
    monkeypatch.delenv('DY_ODOO_ENVIRONMENT',raising=False)
    template=load_template(module)
    material=tmp_path/'签认材料.txt';material.write_text('已确认的脱敏材料')
    bound=[dict(reference=material.name,sha256=hashlib.sha256(material.read_bytes()).hexdigest(),requirements=['按本期清单配置'])]
    q1,q2=template['questions'][:2]
    def call(name,**args):
        result,error=server.execute_tool(name,dict(trace_id='workflow-complete',**args),tmp_path)
        assert error is None,error
        return result
    def save(batch,answers,**scope):
        return call('dy_research_answers_save',module=module,scope='甲院/一期',template_revision=template['revision'],
            batch_id=batch,confirmed=True,answers=answers,**scope)
    first=save('聊天首批',[dict(question_id=q1['question_id'],answer='最初业务回答',research_complete=True,expected_revision=None,attachments=bound)])
    synchronized=Path(first['path']);snapshot=synchronized.read_bytes()
    old=call('dy_research_unanswered_export',module=module,confirmed=True)
    chat=save('聊天修订',[dict(question_id=q1['question_id'],answer='本期最新聊天确认',research_complete=True,
        expected_revision=first['records'][0]['revision'])],scope_id=first['scope_id'])
    book=load_workbook(old['path']);book.active.cell(q2['row'],new_col,'同模板旧答卷的新题回答')
    book.save(old['path']);book.close()
    candidate=call('dy_research_answers_import',path=old['path'])
    assert candidate['can_write_confirmed'] is False and len(candidate['answers'])==1
    assert candidate['answers'][0]['question_id']==q2['question_id']
    assert call('dy_module_memory_read',module=module)['answer_version']==chat['answer_version']
    merged=save('业务确认回填',[{**candidate['answers'][0],'research_complete':True,'attachments':bound}],scope_id=first['scope_id'])
    values={r['question_id']:r for r in merged['records']}
    assert values[q1['question_id']]['answer']=='本期最新聊天确认'
    assert values[q2['question_id']]['answer']=='同模板旧答卷的新题回答'
    assert synchronized.read_bytes()==snapshot
    latest=call('dy_research_unanswered_export',module=module,confirmed=True)
    assert latest['path'] not in {first['path'],old['path']}
    book=load_workbook(latest['path'])
    assert book.active.cell(q1['row'],confirmed_col).value==values[q1['question_id']]['answer']
    assert book.active.cell(q2['row'],confirmed_col).value==values[q2['question_id']]['answer']
    assert all(book.active.cell(q['row'],new_col).value is None for q in template['questions']);book.close()
    prep=call('dy_inspection_prepare',module=module,scope_id=first['scope_id'],environment='UAT')
    assert prep['basis']['answer_version']==merged['answer_version'] and len(prep['expected_checks'])==len(template['questions'])
    report=make_report(tmp_path,module,state=merged)
    validation=call('dy_inspection_validate',report=report)
    assert validation['valid'] and validation['gate']=='待确认'
    artifact=call('dy_inspection_report_save',report=report,confirmed=True)
    content=call('dy_inspection_report_read',module=module,report_id=artifact['report_id'])
    total = len(template['questions'])
    assert f'{total}/{total}' in content and '按本期清单配置' in content
    assert (tmp_path/artifact['pdf_path']).read_bytes().startswith(b'%PDF')
    material.write_text('材料已经修改')
    assert not call('dy_inspection_validate',report=report)['valid']
