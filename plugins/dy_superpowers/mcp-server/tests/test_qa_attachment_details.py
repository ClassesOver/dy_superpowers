"""确认批次内的附件字段明细与同步答卷，不接触真实业务。"""
import copy
import hashlib
import json
import math
from pathlib import Path

import pytest
from openpyxl import load_workbook

import server
from dy_superpowers_mcp.qa_store import commit_qa, export_qa, parse_qa_workbook, read_qa
from dy_superpowers_mcp.qa_templates import load_template


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setenv('CODEX_HOME', str(tmp_path / 'codex'))


def _bound(root, name='材料.txt', *, app='实时附件提取的应用', row=None):
    path = root / name
    path.write_text('原附件包含需提取的配置及无关私密内容，不能自动复制全文', encoding='utf-8')
    return dict(reference=name, sha256=hashlib.sha256(path.read_bytes()).hexdigest(), requirements=['本期应用按确认清单安装'],
        details=[dict(sheet='01_应用安装', rows=[row or {'序号': 1, '应用名称（含英文编码）': app, '是否需要安装': True}])])


def _answer(bound, index=0, revision=None, text='本期启用对应应用'):
    return dict(question_id=load_template('expense')['questions'][index]['question_id'], answer=text,
        research_complete=True, expected_revision=revision, attachments=[bound])


def _save(root, answers, batch='首次', confirmed=True):
    return commit_qa(root, 'expense', scope='甲院/一期', template_revision=load_template('expense')['revision'],
        answers=answers, batch_id=batch, confirmed=confirmed)


def _business_rows(path):
    sheet = load_workbook(path)['01_应用安装']
    return [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2)]


def test_chat_attachment_details_mcp_save_and_both_workbooks_preserve_real_columns(tmp_path):
    bound = _bound(tmp_path)
    request = dict(jsonrpc='2.0', id=1, method='tools/call', params=dict(name='dy_research_answers_save', arguments=dict(
        trace_id='details-save', module='expense', scope='甲院/一期', template_revision=load_template('expense')['revision'],
        answers=[_answer(bound)], batch_id='实时回答附件', confirmed=True)))
    response = server.handle(request, tmp_path)
    assert 'error' not in response, response
    state = json.loads(response['result']['content'][0]['text'])
    assert state['records'][0]['attachments'] == [bound]
    for path in (state['path'], export_qa(tmp_path, 'expense', confirmed=True)['path']):
        row = _business_rows(path)[0]
        assert row[:4] == [1, '实时附件提取的应用', None, True]
        assert row[7:10] == ['应用安装', load_template('expense')['questions'][0]['question'], '材料.txt']
        assert row[10] == bound['requirements'][0]
    _, error = server.execute_tool('dy_module_memory_read', {'trace_id': 'read-attached', 'module': 'expense'}, tmp_path)
    assert error is None
    events = [json.loads(line) for path in tmp_path.rglob('*.jsonl') for line in path.read_text().splitlines()]
    trace = json.dumps(events, ensure_ascii=False)
    assert '实时附件提取的应用' not in trace and '无关私密内容' not in trace
    for event in events:
        if event.get('event') in {'research.write', 'memory.module.read'}:
            assert 'details' not in event.get('research_content', '')


def test_shared_sheet_keeps_each_question_attachment_row_and_candidate_binding(tmp_path):
    one = _bound(tmp_path, '甲.txt', app='甲问题应用')
    one['details'][0]['rows'].append({'序号': 2, '应用名称（含英文编码）': '甲问题第二应用'})
    two = _bound(tmp_path, '乙.txt', app='乙问题应用')
    three = _bound(tmp_path, '丙.txt', app='同题另一附件应用')
    # 同一个业务问题可绑定多份材料；分别追加，不能覆盖先前行。
    answers = [_answer(one)]
    answers[0]['attachments'].append(two)
    answers[0]['attachments'].append(three)
    state = _save(tmp_path, answers)
    exported = export_qa(tmp_path, 'expense', confirmed=True)
    rows = _business_rows(exported['path'])
    assert [r[1] for r in rows] == ['甲问题应用', '甲问题第二应用', '乙问题应用', '同题另一附件应用']
    assert [r[7] for r in rows] == ['应用安装'] * 4
    assert [r[9] for r in rows] == ['甲.txt', '甲.txt', '乙.txt', '丙.txt']
    book = load_workbook(exported['path'])
    book['01_应用安装']['B4'] = '乙问题候选修改'
    book.save(exported['path'])
    parsed = parse_qa_workbook(tmp_path, exported['path'])
    assert not parsed['can_write_confirmed'] and parsed['answers'] == []
    assert parsed['attachment_changes'][0]['question_ids'] == [answers[0]['question_id']]
    assert parsed['attachment_changes'][0]['before'] == '乙问题应用'
    assert read_qa(tmp_path, 'expense')['answer_version'] == state['answer_version']
    assert read_qa(tmp_path, 'expense')['records'][0]['attachments'] == [one, two, three]
    revised = copy.deepcopy(two)
    revised['details'][0]['rows'][0]['应用名称（含英文编码）'] = parsed['attachment_changes'][0]['value']
    updated = _answer(one, revision=state['records'][0]['revision'])
    updated['attachments'] = [one, revised, three]
    confirmed = _save(tmp_path, [updated], '候选人工确认保存')
    assert _business_rows(confirmed['path'])[2][1] == '乙问题候选修改'
    assert _business_rows(state['path'])[2][1] == '乙问题应用'


def test_details_revision_preservation_clearing_and_immutable_workbook_agree(tmp_path):
    bound = _bound(tmp_path)
    first = _save(tmp_path, [_answer(bound)])
    old_path = first['path']
    prior = first['records'][0]
    answer = _answer(bound, revision=prior['revision'], text='仅修订文字')
    del answer['attachments']
    preserved = _save(tmp_path, [answer], '保留附件')
    assert preserved['records'][0]['attachments'] == [bound]
    updated_bound = copy.deepcopy(bound)
    updated_bound['details'][0]['rows'][0]['应用名称（含英文编码）'] = '确认新明细'
    updated = _save(tmp_path, [_answer(updated_bound, revision=preserved['records'][0]['revision'], text='仅修订文字')], '修订明细')
    assert updated['answer_version'] != preserved['answer_version']
    assert updated['records'][0]['revision'] != preserved['records'][0]['revision']
    assert _business_rows(updated['path'])[0][1] == '确认新明细'
    assert _business_rows(old_path)[0][1] == '实时附件提取的应用'
    blank = _answer(updated_bound, revision=updated['records'][0]['revision'], text='仅修订文字')
    blank['attachments'] = []
    cleared = _save(tmp_path, [blank], '清除附件')
    assert cleared['records'][0]['attachments'] == []
    assert _business_rows(cleared['path']) == []


@pytest.mark.parametrize('mutation', ['wrong_sheet', 'unknown_column', 'object', 'nan', 'sensitive', 'unconfirmed'])
def test_details_invalid_or_unconfirmed_batch_never_publishes(tmp_path, mutation):
    bound = _bound(tmp_path)
    if mutation == 'wrong_sheet':
        bound['details'][0]['sheet'] = '02_公司'
    elif mutation == 'unknown_column':
        bound['details'][0]['rows'][0]['未知字段'] = '越界内容'
    elif mutation == 'object':
        bound['details'][0]['rows'][0]['序号'] = {'嵌套': '不允许'}
    elif mutation == 'nan':
        bound['details'][0]['rows'][0]['序号'] = math.nan
    elif mutation == 'sensitive':
        bound['details'][0]['rows'][0]['应用名称（含英文编码）'] = '13812345678'
    with pytest.raises(ValueError):
        _save(tmp_path, [_answer(bound)], confirmed=mutation != 'unconfirmed')
    assert read_qa(tmp_path, 'expense')['records'] == []


def test_legacy_attachment_binding_without_details_stays_unchanged(tmp_path):
    bound = _bound(tmp_path)
    del bound['details']
    state = _save(tmp_path, [_answer(bound)])
    assert state['records'][0]['attachments'] == [bound]
    assert _business_rows(state['path'])[0][:7] == [None] * 7


def test_general_ledger_details_keep_iso_date_json_scalars_and_formula_like_text(tmp_path):
    template = load_template('general-ledger')
    question = next(q for q in template['questions'] if q['check_order'] == 'GL-3-011')
    bound = _bound(tmp_path)
    page = '会计年度与分配公司清单'
    bound['details'] = [{'sheet': page, 'rows': [{'年度开始日期 *': '2026-10-01', '年度结束日期 *': '2027-09-30', '年度编码 *': '=2026'}]}]
    state = commit_qa(tmp_path, 'general-ledger', scope='甲院/一期', template_revision=template['revision'],
        answers=[dict(question_id=question['question_id'], answer='已确认本期会计年度', research_complete=True,
                      expected_revision=None, attachments=[bound])], batch_id='确认年度', confirmed=True)
    sheet = load_workbook(state['path'])[page]
    fields = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(2, fields['年度开始日期 *']).value == '2026-10-01'
    assert sheet.cell(2, fields['年度结束日期 *']).value == '2027-09-30'
    assert sheet.cell(2, fields['年度编码 *']).data_type == 's'
    assert read_qa(tmp_path, 'general-ledger')['records'][0]['attachments'] == [bound]


def test_detail_blank_text_and_finite_scalar_values_have_stable_snapshot(tmp_path):
    row = {'序号': 1.25, '应用名称（含英文编码）': '', '依赖应用': None, '是否需要安装': False}
    bound = _bound(tmp_path, row=row)
    state = _save(tmp_path, [_answer(bound)])
    assert read_qa(tmp_path, 'expense')['records'][0]['attachments'] == [bound]
    assert _business_rows(state['path'])[0][:4] == [1.25, None, None, False]
