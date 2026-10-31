"""业务答卷不展示内部标识，隐藏关联仍支持准确续答。"""
import hashlib
from pathlib import Path

import pytest
from openpyxl import load_workbook

from dy_superpowers_mcp.qa_store import commit_qa, export_qa, parse_qa_workbook
from dy_superpowers_mcp.qa_templates import load_template


def save(root, attachment, *, previous=None, batch='首次'):
    template = load_template('expense')
    question = template['questions'][0]
    return commit_qa(root, 'expense', scope='甲院', template_revision=template['revision'],
        answers=[dict(question_id=question['question_id'], answer='本期使用确认应用', research_complete=True,
                      expected_revision=previous['records'][0]['revision'] if previous else None,
                      attachments=[attachment])], batch_id=batch, confirmed=True)


def bound(root):
    file = root / '材料.txt'
    file.write_text('确认清单', encoding='utf-8')
    return dict(reference=file.name, sha256=hashlib.sha256(file.read_bytes()).hexdigest(), requirements=['使用确认应用'],
        details=[dict(sheet='01_应用安装', rows=[{'应用名称（含英文编码）': '确认应用'}])])


def test_export_hides_internal_identifiers_and_keeps_candidate_binding(tmp_path):
    attachment = bound(tmp_path)
    state = save(tmp_path, attachment)
    path = export_qa(tmp_path, 'expense', confirmed=True)['path']
    book = load_workbook(path)
    hidden = [sheet for sheet in book if sheet.sheet_state == 'veryHidden']
    assert len(hidden) >= 2
    visible = [cell.value for sheet in book if sheet.sheet_state == 'visible' for row in sheet for cell in row]
    assert state['records'][0]['question_id'] not in visible
    assert attachment['sha256'] not in visible
    assert '确认应用' in visible and '材料.txt' in visible
    book['01_应用安装']['B2'] = '候选应用'
    book.save(path)
    candidate = parse_qa_workbook(tmp_path, path)['attachment_changes'][0]
    assert candidate['question_ids'] == [state['records'][0]['question_id']]
    assert candidate['value'] == '候选应用'
    book = load_workbook(path)
    mapping = next(sheet for sheet in book if sheet.title != '_dy_qa' and sheet.sheet_state == 'veryHidden')
    mapping['A2'] = '伪造关联'
    book.save(path)
    with pytest.raises(ValueError):
        parse_qa_workbook(tmp_path, path)


def test_omitted_details_preserve_same_binding_and_explicit_empty_clears(tmp_path):
    attachment = bound(tmp_path)
    first = save(tmp_path, attachment)
    without_details = {key: value for key, value in attachment.items() if key != 'details'}
    preserved = save(tmp_path, without_details, previous=first, batch='仅重新确认附件要求')
    assert preserved['records'][0]['attachments'][0]['details'] == attachment['details']
    assert preserved['answer_version'] == first['answer_version']
    cleared = save(tmp_path, {**without_details, 'details': []}, previous=preserved, batch='明确清空明细')
    assert cleared['records'][0]['attachments'][0]['details'] == []
    assert '确认应用' not in [cell.value for row in load_workbook(cleared['path'])['01_应用安装'] for cell in row]


def test_shared_business_attachment_page_keeps_hidden_question_ownership(tmp_path):
    template = load_template('general-ledger')
    page = next(page for page in template['attachments'] if page['name'] == '报表清单')
    questions = [q for q in template['questions'] if '报表清单' in q['attachment_names']][:2]
    assert len(questions) == 2
    attachment = bound(tmp_path)
    answers = []
    for number, question in enumerate(questions, 1):
        details = [{'sheet': '报表清单', 'rows': [{page['headers'][1]: f'业务报表{number}'}]}]
        answers.append(dict(question_id=question['question_id'], answer=f'确认报表{number}', research_complete=True,
                            expected_revision=None, attachments=[{**attachment, 'details': details}]))
    state = commit_qa(tmp_path, 'general-ledger', scope='甲院', template_revision=template['revision'],
                      answers=answers, batch_id='多问题材料', confirmed=True)
    path = export_qa(tmp_path, 'general-ledger', confirmed=True)['path']
    book = load_workbook(path)
    book['报表清单']['B3'] = '第二问题修订'
    book.save(path)
    candidate = parse_qa_workbook(tmp_path, path)['attachment_changes'][0]
    assert candidate['question_ids'] == [questions[1]['question_id']]
    assert candidate['before'] == '业务报表2'
    assert state['confirmed_count'] == 2


@pytest.mark.parametrize('module', ['expense', 'general-ledger'])
def test_answer_columns_merge_all_rows_of_each_projected_question(tmp_path, module):
    template = load_template(module)
    exported = export_qa(tmp_path, module, confirmed=True)
    sheet = load_workbook(exported['path']).worksheets[0]
    ranges = set(map(str, sheet.merged_cells.ranges))
    for question in template['questions']:
        rows = question['rows']
        if len(rows) > 1:
            for column in 'FGH':
                assert f'{column}{rows[0]}:{column}{rows[-1]}' in ranges
