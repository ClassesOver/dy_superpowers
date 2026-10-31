"""新模板 QA 的 MCP 与巡检边界；仅使用临时工作区。"""
import copy
import hashlib
import sys
from datetime import datetime, timezone
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parents[1] / 'src'))
import server
from dy_superpowers_mcp.qa_store import commit_qa, read_qa
from dy_superpowers_mcp.qa_templates import load_template
from dy_superpowers_mcp.qa_workbook import answer_column
from dy_superpowers_mcp.trace import TraceStore


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setenv('CODEX_HOME', str(tmp_path / 'codex'))
    monkeypatch.setattr(server, 'initialize_telemetry', lambda: False)
    monkeypatch.delenv('DB_NAME', raising=False)
    monkeypatch.delenv('DY_ODOO_ENVIRONMENT', raising=False)


def call(root, name, **args):
    result, error = server.execute_tool(name, {'trace_id': 'qa-integration', **args}, root)
    assert error is None, error
    return result


def saved(root, module='expense', count=1, attachment=True):
    template = load_template(module)
    (root / 'evidence.txt').write_text('经签认的配置核对材料', encoding='utf-8')
    bound = [{'reference': 'evidence.txt', 'sha256': hashlib.sha256((root / 'evidence.txt').read_bytes()).hexdigest(),
              'requirements': ['应用安装符合清单', '各公司覆盖完整']}] if attachment else []
    return commit_qa(root, module, scope='甲院/一期', template_revision=template['revision'],
                     answers=[dict(question_id=q['question_id'], answer='本期按确认清单配置', research_complete=True,
                                   expected_revision=None, attachments=bound) for q in template['questions'][:count]],
                     batch_id='batch-first', confirmed=True, confirmed_by='测试签认')


def report_for(root, module='expense', count=1):
    state = saved(root, module, count)
    verification = root / 'independent-verification.txt'
    verification.write_text('独立核验系统现状的材料', encoding='utf-8')
    basis = {k: state[k] for k in ('module', 'scope_id', 'template_revision', 'answer_version')}
    basis['environment'] = 'UAT'
    records = []
    answers = {r['question_id']: r for r in state['records']}
    for question in state['questions']:
        qid = question['question_id']
        answer = answers.get(qid)
        rec = dict(question_id=qid, answer_revision=answer['revision'] if answer else None,
                   verification_method='材料核对', verified_at=datetime.now(timezone.utc).isoformat(),
                   reason='完整材料逐项核对' if answer else '尚无确认答案',
                   requirements=[], evidence_refs=[], result='通过' if answer else '未检查')
        if answer:
            attachment = answer['attachments'][0]
            evidence = dict(kind='document', reference=verification.name, sha256=hashlib.sha256(verification.read_bytes()).hexdigest(),
                            observed_at=datetime.now(timezone.utc).isoformat())
            rec['evidence_refs'] = [evidence]
            rec['requirements'] = [dict(requirement_id=f'{qid}-{index}', text=text,
                source_refs=[dict(kind='answer', question_id=qid, revision=answer['revision']),
                             dict(kind='attachment', reference=attachment['reference'], sha256=attachment['sha256'])],
                evidence_refs=[evidence], verification_target='document', result='通过') for index, text in enumerate(attachment['requirements'])]
            rec['verified_at'] = datetime.now(timezone.utc).isoformat()
        records.append(rec)
    return dict(basis=basis, expected_checks=[dict(question_id=q['question_id']) for q in state['questions']], records=records)


def validate(root, report):
    return call(root, 'dy_inspection_validate', report=report)


@pytest.mark.parametrize('module', ['expense', 'general-ledger'])
def test_new_tools_batch_save_read_export_prepare(tmp_path, module):
    template = load_template(module)
    assert 'dy_research_answers_save' in server.TOOL_SCHEMAS
    assert 'dy_research_answers_import' in server.TOOL_SCHEMAS
    assert 'confirmed_by' not in server.TOOL_SCHEMAS['dy_research_answers_save']['properties']
    result = call(tmp_path, 'dy_research_answers_save', module=module, scope='甲院/一期',
                  template_revision=template['revision'], batch_id='chat-first', confirmed=True,
                  answers=[dict(question_id=template['questions'][0]['question_id'], answer='本期启用',
                                research_complete=False, expected_revision=None)])
    assert (tmp_path / result['path']).is_file()
    state = call(tmp_path, 'dy_module_memory_read', module=module)
    assert state['confirmed_count'] == 1 and state['answered'] == 0
    progress = call(tmp_path, 'dy_research_progress', module=module, intent='research')
    assert progress['total'] == len(template['questions']) and progress['next_action'] == 'offer_checklist'
    exported = call(tmp_path, 'dy_research_unanswered_export', module=module, confirmed=True)
    assert exported['path'] != result['path']
    parsed = call(tmp_path, 'dy_research_answers_import', path=exported['path'])
    assert parsed['can_write_confirmed'] is False
    prep = call(tmp_path, 'dy_inspection_prepare', module=module, scope_id=state['scope_id'], environment='UAT')
    assert len(prep['items']) == len(template['questions'])
    assert prep['items'][0]['readiness'] == 'incomplete'
    assert prep['basis']['answer_version'] == state['answer_version']


@pytest.mark.parametrize('module', ['expense', 'general-ledger'])
def test_inspection_and_history_do_not_read_static_matrix(tmp_path, monkeypatch, module):
    original = Path.read_bytes

    def without_matrix(path):
        if path.name == 'inspection-matrix.md':
            raise FileNotFoundError(path)
        return original(path)

    monkeypatch.setattr(Path, 'read_bytes', without_matrix)
    report = report_for(tmp_path, module)
    assert validate(tmp_path, report)['valid'] is True
    result = call(tmp_path, 'dy_inspection_report_save', report=report, confirmed=True)
    assert (tmp_path / result['path']).is_file()
    from dy_superpowers_mcp.research import write_record
    write_record(tmp_path, dict(module=module, check_order='historical', research_id='01',
        config_item='历史记录', conclusion='仅供历史回查，不作为当前巡检依据', scope='历史范围',
        confirmed=True, confirmed_by='测试', confirmed_at='2026-10-30'))
    assert call(tmp_path, 'dy_module_memory_read', module=module)['confirmed_count'] == 1


def test_legacy_answers_never_enter_default_read_or_progress(tmp_path):
    from dy_superpowers_mcp.research import write_record
    write_record(tmp_path, dict(module='expense', check_order='1', research_id='01', config_item='历史',
        conclusion='历史答案', scope='历史范围', confirmed=True, research_complete=True, confirmed_by='测试', confirmed_at='2026-10-30'))
    assert call(tmp_path, 'dy_module_memory_read', module='expense')['records'] == []
    assert call(tmp_path, 'dy_research_progress', module='expense', intent='progress')['answered'] == 0
    assert call(tmp_path, 'dy_module_memory_read', module='expense', legacy=True)['records']


@pytest.mark.parametrize('module', ['expense', 'general-ledger'])
def test_template_report_keeps_all_current_questions_and_waiting_gate(tmp_path, module):
    report = report_for(tmp_path, module)
    result = validate(tmp_path, report)
    assert result['valid'] is True, result['errors']
    assert result['gate'] == '待确认'
    assert next(iter(result['coverage'].values()))['template'] == len(load_template(module)['questions'])
    assert not validate(tmp_path, {'expected_checks': [], 'records': []})['valid']


@pytest.mark.parametrize('mutate', [
    lambda r: r['records'].pop(),
    lambda r: r['expected_checks'].pop(),
    lambda r: r['records'][0]['requirements'].pop(),
    lambda r: r['records'][0].update(requirements=[]),
    lambda r: r['basis'].update(answer_version='stale'),
    lambda r: r['records'][0].update(answer_revision='stale'),
    lambda r: r['records'][0]['requirements'][0]['source_refs'][0].update(question_id='other'),
    lambda r: r['records'][0]['requirements'][0]['evidence_refs'][0].update(reference='fake.txt'),
    lambda r: r['records'][0].update(verified_at='2020-01-01T00:00:00+00:00'),
    lambda r: r['records'][0]['requirements'][0].update(result='待确认'),
])
def test_report_rejects_missing_stale_and_forged_boundaries(tmp_path, mutate):
    report = report_for(tmp_path)
    mutate(report)
    assert validate(tmp_path, report)['valid'] is False


def test_attachment_change_invalidates_report(tmp_path):
    report = report_for(tmp_path)
    (tmp_path / 'evidence.txt').write_text('材料已变化', encoding='utf-8')
    assert validate(tmp_path, report)['valid'] is False


@pytest.fixture
def configured_inspection(monkeypatch):
    for key, value in dict(APP_URL='https://odoo.example.test', DB_NAME='UAT', APP_USER='operator', APP_PASSWORD='placeholder').items():
        monkeypatch.setenv(key, value)
    from dy_superpowers_mcp.qa_basis import configured_environment_info
    return configured_environment_info()['connection_fingerprint']


def test_successful_odoo_trace_requires_full_current_context(tmp_path, configured_inspection):
    from dy_superpowers_mcp.odoo_comparison import compare_records
    comparison = {'expected_count': 1, 'expected_values': {'state': 'installed'}}
    report = report_for(tmp_path)
    report['basis']['connection_fingerprint'] = configured_inspection
    record = report['records'][0]
    evidence = dict(kind='odoo', reference='odoo://ir.module.module', trace_id='odoo-proof', model='ir.module.module', comparison=comparison)
    record['evidence_refs'] = [evidence]
    for req in record['requirements']:
        req['verification_target'] = 'configuration'
        req['evidence_refs'] = [evidence]
    assert validate(tmp_path, report)['valid'] is False
    context = {**report['basis'], 'question_id': record['question_id'], 'answer_revision': record['answer_revision']}
    store = TraceStore(tmp_path)
    store.append('odoo-proof', dict(event='odoo.read', status='ok', tool='dy_odoo_record_search',
        operation='search_read', model='ir.module.module', span_id='proof-1', trace_context=context,
        **compare_records(comparison, [{'state': 'installed'}])))
    store.append('odoo-proof', dict(event='tool.end', status='ok', tool='dy_odoo_record_search', span_id='proof-1'))
    record['verified_at'] = datetime.now(timezone.utc).isoformat()
    assert validate(tmp_path, report)['valid'] is True
    report['basis']['environment'] = 'PROD'
    assert validate(tmp_path, report)['valid'] is False

def test_report_save_read_and_pdf_use_template_business_requirements(tmp_path):
    report = report_for(tmp_path)
    result = call(tmp_path, 'dy_inspection_report_save', report=report, confirmed=True)
    content = call(tmp_path, 'dy_inspection_report_read', module='expense', report_id=result['report_id'])
    assert '应用安装符合清单' in content
    assert '范围：甲院/一期' in content
    assert (tmp_path / result['pdf_path']).read_bytes().startswith(b'%PDF')


def test_skill_contracts_use_template_batches_and_do_not_require_old_matrix():
    root = Path(__file__).parents[2] / 'skills'
    research = (root / 'module-research/SKILL.md').read_text()
    assert 'dy_research_answers_save' in research
    assert 'dy_research_answers_import' in research
    assert 'dy_research_record_write(' not in research
    for module in ('expense', 'general-ledger'):
        text = (root / f'{module}-inspection/SKILL.md').read_text()
        assert 'question_id' in text
        assert '必读完整矩阵' not in text

@pytest.mark.parametrize('mutate', [
    lambda r: r['records'][0].update(result=[]),
    lambda r: r['records'][0]['requirements'][0]['source_refs'][1].update(reference={}),
    lambda r: r['basis'].update(environment='中文环境'),
])
def test_malformed_report_is_invalid_without_crash(tmp_path, mutate):
    report = report_for(tmp_path)
    mutate(report)
    result = validate(tmp_path, report)
    assert result['valid'] is False


def test_definitive_requirement_cannot_hide_under_waiting_question(tmp_path):
    report = report_for(tmp_path)
    current = read_qa(tmp_path, 'expense')
    first = current['records'][0]
    updated = commit_qa(tmp_path, 'expense', scope=current['scope'], scope_id=current['scope_id'],
        template_revision=current['template_revision'], batch_id='partial-again', confirmed=True,
        answers=[dict(question_id=first['question_id'], answer=first['answer'], research_complete=False,
                      expected_revision=first['revision'])])
    report['basis']['answer_version'] = updated['answer_version']
    record = report['records'][0]
    record['result'] = '待确认'
    record['answer_revision'] = updated['records'][0]['revision']
    for req in record['requirements']:
        req['source_refs'][0]['revision'] = record['answer_revision']
    assert validate(tmp_path, report)['valid'] is False

def test_real_stdio_batch_cross_session_workbook_and_report(tmp_path):
    import asyncio
    import json
    import os
    from mcp import ClientSession, StdioServerParameters
    from mcp.client.stdio import stdio_client
    from openpyxl import load_workbook
    script = Path(__file__).parents[1] / 'src/server.py'
    env = {**os.environ, 'CODEX_HOME': str(tmp_path / 'codex'), 'DB_NAME': 'stdio-db',
           'DY_ODOO_ENVIRONMENT': 'UAT', 'APP_URL': 'https://stdio.invalid',
           'APP_USER': 'stdio-user-secret', 'APP_PASSWORD': 'stdio-password-secret'}
    params = StdioServerParameters(command=sys.executable, args=[str(script)], cwd=str(tmp_path), env=env)
    async def invoke(session, name, **args):
        result = await session.call_tool(name, {'trace_id': 'stdio-qa-new', **args})
        assert not result.is_error, result
        return json.loads(result.content[0].text)
    async def first_session():
        async with stdio_client(params) as (reader, writer):
            async with ClientSession(reader, writer) as session:
                await session.initialize()
                schemas = {t.name: t.input_schema for t in (await session.list_tools()).tools}
                assert {'dy_research_answers_save', 'dy_research_answers_import'} <= schemas.keys()
                assert 'dy_odoo_environment_info' in schemas
                identity = await invoke(session, 'dy_odoo_environment_info')
                assert {key: identity[key] for key in ('environment', 'database', 'configured')} == {'environment': 'UAT', 'database': 'stdio-db', 'configured': True}
                assert len(identity['connection_fingerprint']) == 64
                template = load_template('expense')
                return await invoke(session, 'dy_research_answers_save', module='expense', scope='甲院/一期',
                    template_revision=template['revision'], batch_id='stdio-initial', confirmed=True,
                    answers=[dict(question_id=template['questions'][0]['question_id'], answer='已确认局部回答',
                                  research_complete=False, expected_revision=None)])
    state = asyncio.run(first_session())
    async def resumed_session():
        async with stdio_client(params) as (reader, writer):
            async with ClientSession(reader, writer) as session:
                await session.initialize()
                latest = await invoke(session, 'dy_module_memory_read', module='expense')
                assert latest['answer_version'] == state['answer_version']
                output = await invoke(session, 'dy_research_unanswered_export', module='expense', confirmed=True)
                book = load_workbook(output['path'])
                book.worksheets[0].cell(load_template('expense')['questions'][1]['row'], answer_column(load_template('expense')) + 1).value = '本轮新增回答'
                book.save(output['path'])
                parsed = await invoke(session, 'dy_research_answers_import', path=output['path'])
                assert parsed['answers'][0]['research_complete'] is False
                updated = await invoke(session, 'dy_research_answers_save', module='expense', scope=latest['scope'],
                    scope_id=latest['scope_id'], template_revision=latest['template_revision'],
                    batch_id='stdio-update', confirmed=True, answers=parsed['answers'])
                assert updated['confirmed_count'] == 2
                prep = await invoke(session, 'dy_inspection_prepare', module='expense', scope_id=latest['scope_id'])
                assert prep['basis']['environment'] == 'UAT'
                answers = {a['question_id']: a for a in updated['records']}
                report = dict(basis=prep['basis'], expected_checks=prep['expected_checks'], records=[dict(
                    question_id=q['question_id'], answer_revision=answers.get(q['question_id'], {}).get('revision'),
                    verification_method='尚未实时验证', verified_at=datetime.now(timezone.utc).isoformat(),
                    reason='等待完整确认或业务证据', requirements=[], evidence_refs=[], result='待确认') for q in prep['items']])
                checked = await invoke(session, 'dy_inspection_validate', report=report)
                assert checked['valid'] and checked['gate'] == '待确认'
                saved_report = await invoke(session, 'dy_inspection_report_save', report=report, confirmed=True)
                result = await session.call_tool('dy_inspection_report_read', dict(trace_id='stdio-qa-new', module='expense', report_id=saved_report['report_id']))
                total = len(load_template('expense')['questions'])
                assert f'{total}/{total}' in result.content[0].text
    asyncio.run(resumed_session())

@pytest.mark.parametrize('module', ['expense', 'general-ledger'])
def test_all_current_points_can_be_not_applicable_with_real_basis(tmp_path, module):
    report = report_for(tmp_path, module, len(load_template(module)['questions']))
    for record in report['records']:
        record['result'] = '不适用'
        record['reason'] = '本期已确认该项不适用，材料签认范围明确'
        for requirement in record['requirements']:
            requirement['result'] = '不适用'
    checked = validate(tmp_path, report)
    assert checked['valid'], checked['errors']
    assert checked['gate'] == '不适用'
    report['records'][0]['result'] = '不通过'
    report['records'][0]['requirements'][0]['result'] = '不通过'
    assert validate(tmp_path, report)['gate'] == '不通过'


@pytest.mark.parametrize('count, route', [(17, 'excel'), (18, 'interactive')])
def test_new_mcp_progress_keeps_70_percent_and_explicit_export(tmp_path, count, route):
    saved(tmp_path, count=count, attachment=False)
    progress = call(tmp_path, 'dy_research_progress', module='expense', intent='research')
    assert progress['route'] == route
    assert progress['answered'] == count and progress['total'] == len(load_template('expense')['questions'])
    assert call(tmp_path, 'dy_research_unanswered_export', module='expense', confirmed=True)['path']


def test_read_multiple_scopes_keeps_progress_independent(tmp_path):
    first = saved(tmp_path, attachment=False)
    template = load_template('expense')
    commit_qa(tmp_path, 'expense', scope='乙院/二期', template_revision=template['revision'],
        batch_id='other-scope', confirmed=True,
        answers=[dict(question_id=template['questions'][1]['question_id'], answer='另院局部',
                      research_complete=False, expected_revision=None)])
    progress = call(tmp_path, 'dy_research_progress', module='expense', intent='progress')
    assert progress['scope_mode'] == 'multiple'
    assert progress['records'] == []
    assert sorted(s['answered'] for s in progress['scope_progress']) == [0, 1]
    selected = call(tmp_path, 'dy_module_memory_read', module='expense', scope_id=first['scope_id'])
    assert selected['answered'] == 1

@pytest.mark.parametrize('filters', [dict(scope='其他范围'), dict(check_order='1', research_id='01'), dict(limit=1), dict(checks=[dict(check_order='1', research_id='01')])])
def test_new_memory_view_rejects_legacy_filters_instead_of_ignoring(tmp_path, filters):
    result, error = server.execute_tool('dy_module_memory_read', dict(trace_id='filter-boundary', module='expense', **filters), tmp_path)
    assert result is None and error['code'] == -32602
    assert 'legacy' in error['message']


def test_manual_signature_and_odoo_evidence_cannot_postdate_verification(tmp_path):
    report = report_for(tmp_path)
    record = report['records'][0]
    evidence = {**record['evidence_refs'][0], 'kind': 'manual', 'signed_by': '业务签认',
                'signed_at': '2099-01-01T00:00:00+00:00'}
    record['evidence_refs'] = [evidence]
    for requirement in record['requirements']:
        requirement['evidence_refs'] = [evidence]
    assert validate(tmp_path, report)['valid'] is False
    evidence['signed_at'] = '2020-01-01T00:00:00+00:00'
    assert validate(tmp_path, report)['valid'] is False
    evidence['signed_at'] = datetime.now(timezone.utc).isoformat()
    assert validate(tmp_path, report)['valid'] is False
    record['verified_at'] = datetime.now(timezone.utc).isoformat()
    assert validate(tmp_path, report)['valid'] is True


def test_legacy_report_is_history_readable_but_cannot_get_new_pdf(tmp_path):
    from dy_superpowers_mcp.reports import report_path
    path = report_path(tmp_path, 'expense', 'old-report')
    path.parent.mkdir(parents=True)
    path.write_text('# 旧巡检报告\n\n```json\n{"expected_checks": [], "records": []}\n```\n', encoding='utf-8')
    assert '旧巡检报告' in call(tmp_path, 'dy_inspection_report_read', module='expense', report_id='old-report')
    result, error = server.execute_tool('dy_inspection_report_export_pdf', dict(
        trace_id='legacy-pdf', module='expense', report_id='old-report'), tmp_path)
    assert result is None and error['code'] == -32602


def test_environment_info_and_prepare_reuse_configured_identity_without_login(tmp_path, monkeypatch):
    import json
    values = {'APP_URL': 'https://odoo.example.test', 'APP_USER': 'secret-user',
              'APP_PASSWORD': 'secret-password', 'DB_NAME': 'configured-db'}
    for key, value in values.items():
        monkeypatch.setenv(key, value)
    monkeypatch.delenv('DY_ODOO_ENVIRONMENT', raising=False)
    monkeypatch.setattr(server, '_odoo_client', lambda *_: pytest.fail('身份读取不能登录或请求Odoo'))
    assert 'dy_odoo_environment_info' in server.TOOL_SCHEMAS
    info = call(tmp_path, 'dy_odoo_environment_info')
    assert {key: info[key] for key in ('environment', 'database', 'configured')} == {'environment': 'configured-db', 'database': 'configured-db', 'configured': True}
    assert len(info['connection_fingerprint']) == 64
    serialized = json.dumps(info)
    assert all(value not in serialized for key, value in values.items() if key != 'DB_NAME')
    state = saved(tmp_path)
    prepared = call(tmp_path, 'dy_inspection_prepare', module='expense', scope_id=state['scope_id'])
    assert prepared['basis']['environment'] == info['environment']
    monkeypatch.setenv('DY_ODOO_ENVIRONMENT', 'UAT-explicit')
    assert call(tmp_path, 'dy_odoo_environment_info')['environment'] == 'UAT-explicit'


def test_environment_info_missing_configuration_never_guesses_dev_or_prod(tmp_path, monkeypatch):
    for key in ('APP_URL', 'APP_USER', 'APP_PASSWORD', 'DB_NAME', 'DY_ODOO_ENVIRONMENT'):
        monkeypatch.delenv(key, raising=False)
    assert 'dy_odoo_environment_info' in server.TOOL_SCHEMAS
    assert call(tmp_path, 'dy_odoo_environment_info') == {'environment': None, 'database': None, 'configured': False, 'connection_fingerprint': None}
    state = saved(tmp_path)
    result, error = server.execute_tool('dy_inspection_prepare', dict(trace_id='identity-missing',
        module='expense', scope_id=state['scope_id']), tmp_path)
    assert result is None and error['code'] == -32602


def test_environment_override_is_forwarded_and_invalid_value_never_falls_back(tmp_path, monkeypatch):
    import json
    config = json.loads((Path(__file__).parents[2] / '.mcp.json').read_text())
    assert 'DY_ODOO_ENVIRONMENT' in config['mcpServers']['dy-superpowers-odoo']['env_vars']
    monkeypatch.setenv('DB_NAME', 'real-database')
    monkeypatch.setenv('DY_ODOO_ENVIRONMENT', 'invalid environment label')
    info = call(tmp_path, 'dy_odoo_environment_info')
    assert info['environment'] is None and info['database'] == 'real-database'


def test_waiting_report_must_display_all_current_attachment_requirements(tmp_path):
    report = report_for(tmp_path)
    record = report['records'][0]
    record.update(result='待确认', evidence_refs=[])
    for requirement in record['requirements']:
        requirement.update(result='未检查', evidence_refs=[])
    assert validate(tmp_path, report)['valid'] is True
    delivered = call(tmp_path, 'dy_inspection_report_save', report=report, confirmed=True)
    content = call(tmp_path, 'dy_inspection_report_read', module='expense', report_id=delivered['report_id'])
    assert all(text in content for text in ('应用安装符合清单', '各公司覆盖完整', 'evidence'))
    assert (tmp_path / delivered['pdf_path']).read_bytes().startswith(b'%PDF')
    record['requirements'].pop()
    assert validate(tmp_path, report)['valid'] is False
    record['requirements'] = []
    assert validate(tmp_path, report)['valid'] is False


def test_question_odoo_evidence_accepts_only_own_requirement_ids(tmp_path, configured_inspection):
    report = report_for(tmp_path)
    report['basis']['connection_fingerprint'] = configured_inspection
    record = report['records'][0]
    evidence = dict(kind='odoo', reference='odoo://ir.module.module', trace_id='requirement-proof', model='ir.module.module')
    record['evidence_refs'] = [evidence]
    context = {**report['basis'], 'question_id': record['question_id'],
               'answer_revision': record['answer_revision'], 'requirement_id': record['requirements'][0]['requirement_id']}
    store = TraceStore(tmp_path)
    store.append('requirement-proof', dict(event='odoo.read', status='ok', tool='dy_odoo_record_search',
        operation='search_read', model='ir.module.module', span_id='requirement-1', trace_context=context))
    store.append('requirement-proof', dict(event='tool.end', status='ok', tool='dy_odoo_record_search', span_id='requirement-1'))
    record['verified_at'] = datetime.now(timezone.utc).isoformat()
    checked = validate(tmp_path, report)
    assert checked['valid'], checked['errors']
    record['requirements'][0]['requirement_id'] = 'unrelated-requirement'
    assert validate(tmp_path, report)['valid'] is False


def test_question_evidence_can_aggregate_multiple_precise_requirement_traces(tmp_path, configured_inspection):
    from dy_superpowers_mcp.odoo_comparison import compare_records
    comparison = {'expected_count': 1, 'expected_values': {'state': 'installed'}}
    report = report_for(tmp_path)
    report['basis']['connection_fingerprint'] = configured_inspection
    record = report['records'][0]
    store = TraceStore(tmp_path)
    refs = []
    for number, requirement in enumerate(record['requirements']):
        requirement['verification_target'] = 'configuration'
        trace_id, span_id = f'precise-proof-{number}', f'precise-span-{number}'
        evidence = dict(kind='odoo', reference='odoo://ir.module.module', trace_id=trace_id, model='ir.module.module', comparison=comparison)
        context = {**report['basis'], 'question_id': record['question_id'],
                   'answer_revision': record['answer_revision'], 'requirement_id': requirement['requirement_id']}
        store.append(trace_id, dict(event='odoo.read', status='ok', tool='dy_odoo_record_search',
            operation='search_read', model='ir.module.module', span_id=span_id, trace_context=context,
            **compare_records(comparison, [{'state': 'installed'}])))
        store.append(trace_id, dict(event='tool.end', status='ok', tool='dy_odoo_record_search', span_id=span_id))
        requirement['evidence_refs'] = [evidence]
        refs.append(evidence)
    record.update(evidence_refs=refs, verified_at=datetime.now(timezone.utc).isoformat())
    checked = validate(tmp_path, report)
    assert checked['valid'], checked['errors']
    record['requirements'][0]['evidence_refs'] = [refs[1]]
    assert validate(tmp_path, report)['valid'] is False


def test_business_evidence_wrong_type_returns_invalid_without_crash(tmp_path):
    report = report_for(tmp_path)
    requirement = report['records'][0]['requirements'][0]
    requirement.update(verification_target='business', evidence_refs=7)
    assert validate(tmp_path, report)['valid'] is False


def test_report_save_malformed_basis_returns_invalid_arguments(tmp_path):
    result, error = server.execute_tool('dy_inspection_report_save', dict(trace_id='malformed-save',
        report={'basis': [], 'expected_checks': [], 'records': []}, confirmed=True), tmp_path)
    assert result is None and error['code'] == -32602


def test_configured_environment_rejects_forged_labels_before_remote_call(tmp_path, monkeypatch):
    monkeypatch.setenv('DB_NAME', 'dev-real')
    monkeypatch.delenv('DY_ODOO_ENVIRONMENT', raising=False)
    monkeypatch.setattr(server, '_odoo_client', lambda *_: pytest.fail('伪标签拒绝必须发生在远程调用前'))
    state = saved(tmp_path)
    for tool, arguments in [
        ('dy_inspection_prepare', dict(module='expense', scope_id=state['scope_id'], environment='PROD')),
        ('dy_odoo_record_search', dict(model='ir.module.module', args=[[]], kwargs={'fields': ['name'], 'limit': 1},
                                       trace_context={'environment': 'PROD'})),
        ('dy_odoo_module_info', dict(module='base', trace_context={'environment': 'PROD'})),
    ]:
        result, error = server.execute_tool(tool, dict(trace_id='forged-environment', **arguments), tmp_path)
        assert result is None and error['code'] == -32602
    report = report_for(tmp_path)
    assert validate(tmp_path, report)['valid'] is False


@pytest.mark.parametrize('value', [None, [], 7, {}])
def test_malformed_nested_collections_never_crash(tmp_path, value):
    from dy_superpowers_mcp.inspection import validate_inspection
    report = report_for(tmp_path)
    for field in ('requirements', 'evidence_refs'):
        broken = copy.deepcopy(report)
        broken['records'][0][field] = value
        assert validate_inspection(tmp_path, broken)['valid'] is False
    for field in ('source_refs', 'evidence_refs'):
        broken = copy.deepcopy(report)
        broken['records'][0]['requirements'][0][field] = value
        assert validate_inspection(tmp_path, broken)['valid'] is False
    for basis in (value, {'module': value}):
        broken = dict(report, basis=basis)
        result, error = server.execute_tool('dy_inspection_report_save', dict(
            trace_id='malformed-types', report=broken, confirmed=True), tmp_path)
        assert result is None and error['code'] == -32602


@pytest.mark.parametrize('basis', [None, [], 7, {}, {'module': []}, {'module': {}}])
def test_pdf_export_malformed_basis_is_explained_invalid_request(tmp_path, basis):
    import json
    from dy_superpowers_mcp.reports import report_path
    path = report_path(tmp_path, 'expense', 'malformed-pdf')
    path.parent.mkdir(parents=True)
    payload = json.dumps(dict(basis=basis, expected_checks=[], records=[]))
    path.write_text('# 历史数据\n\n```json\n' + payload + '\n```\n', encoding='utf-8')
    result, error = server.execute_tool('dy_inspection_report_export_pdf', dict(trace_id='pdf-malformed',
        module='expense', report_id='malformed-pdf'), tmp_path)
    assert result is None and error['code'] == -32602


def test_odoo_environment_trace_is_server_bound_and_unlabeled_reads_remain_available(tmp_path, monkeypatch):
    from types import SimpleNamespace
    monkeypatch.setenv('DB_NAME', 'actual-dev')
    monkeypatch.setattr(server, '_odoo_client', lambda *_: SimpleNamespace(search_read=lambda *args, **kwargs: []))
    assert call(tmp_path, 'dy_odoo_record_search', model='ir.module.module', args=[[]], kwargs={'fields': ['name'], 'limit': 1}) == []
    events = TraceStore(tmp_path).read('qa-integration')
    read = next(e for e in events if e.get('event') == 'odoo.read')
    assert read['trace_context']['environment'] == 'actual-dev'
    monkeypatch.delenv('DB_NAME')
    assert call(tmp_path, 'dy_odoo_record_search', model='ir.module.module', args=[[]], kwargs={'fields': ['name'], 'limit': 1}) == []
