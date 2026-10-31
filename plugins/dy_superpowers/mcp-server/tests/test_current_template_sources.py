"""原始问卷按当前问题分组解析，不依赖旧逐点列。"""
import importlib.util
import json
from datetime import timedelta
from pathlib import Path
import subprocess
import sys

import pytest
from openpyxl import load_workbook
from dy_superpowers_mcp.source_archive import archive_source

PLUGIN = Path(__file__).parents[2]
SCRIPT = PLUGIN / 'skills/module-research/scripts/research_sources.py'
spec = importlib.util.spec_from_file_location('current_template_sources', SCRIPT)
sources = importlib.util.module_from_spec(spec)
spec.loader.exec_module(sources)


@pytest.mark.parametrize('module,name,total', [
    ('expense', '费用模块调研清单.xlsx', 25),
    ('general-ledger', '总账模块调研模板.xlsx', 37),
])
def test_original_template_answers_grouped_and_blanks_counted(tmp_path, module, name, total):
    path = tmp_path / name
    book = load_workbook(PLUGIN / 'skills/module-research/references' / name)
    main = book.worksheets[0]
    main['F3'] = '甲院本期需要基础应用'
    main['F4'] = '仅一期；移动报销不交付'
    book.save(path)
    result = sources.parse_source(path, module)
    assert len(result['records']) == 1
    row = result['records'][0]
    assert row['question_id'] == main['A3'].value
    assert row['conclusion'] == '甲院本期需要基础应用\n仅一期；移动报销不交付'
    assert len(row['answer_cells']) == 2
    assert len(row['attachment_branches']) >= 2
    assert len(result['unanswered']) == total - 1
    assert result['coverage']['question_groups'] == total
    assert main.title not in result['attachment_sheets']
    page = book.worksheets[1]
    header_row = 2 if module == 'expense' else 1
    expected = [str(c.value).strip() if c.value is not None else '' for c in page[header_row]]
    assert result['attachment_sheets'][page.title]['headers'] == expected
    if module == 'general-ledger':
        assert {q['question_id'] for q in result['placeholder_questions']} == {'GL-2-009', 'GL-3-012'}


def test_excel_duration_material_can_be_archived_by_cli(tmp_path):
    path = tmp_path / '时长材料.xlsx'
    book = load_workbook(PLUGIN / 'skills/module-research/references/总账模块调研模板.xlsx')
    page = book.worksheets[1]
    page['A10'] = timedelta(hours=36)
    page['A10'].number_format = '[hh]:mm:ss'
    book.save(path)
    result = sources.parse_source(path, 'general-ledger')
    assert result['attachment_sheets'][page.title]['rows'][-1][0] == 'P1DT12H0M0S'
    json.dumps(result)
    completed = subprocess.run([sys.executable, str(SCRIPT), str(path), '--module', 'general-ledger',
                                '--workspace', str(tmp_path)], capture_output=True, text=True)
    assert completed.returncode == 0, completed.stderr
    output = json.loads(completed.stdout)
    assert (tmp_path / output['files']['candidates']).is_file()


@pytest.mark.parametrize('value,expected', [
    (timedelta(0), 'P0DT0H0M0S'),
    (-timedelta(microseconds=1), '-P0DT0H0M0.000001S'),
])
def test_material_duration_preserves_zero_and_signed_fraction(value, expected):
    assert sources._json_cell(value) == expected


def test_verified_archive_reference_keeps_hash_but_body_stays_redacted(tmp_path):
    digest = 'a' * 10 + '13800138000' + 'b' * 43
    relative = f'.dy_superpowers-parsed/{digest}/runs/20261030T071052435048Z-654f9761'
    target = tmp_path / relative
    target.mkdir(parents=True)
    (target / 'manifest.json').write_text(json.dumps({'source_sha256': digest}))
    source = tmp_path / '材料.txt'
    source.write_text('业务材料')
    result = archive_source(tmp_path, source, 'expense', candidates={
        'records': [{'detail_refs': [{'archive_path': relative, 'source_sha256': digest}],
                     'conclusion': relative + ' 联系人13800138000'}],
        'untrusted': {'archive_path': relative + '/13800138000', 'source_sha256': digest}})
    saved = json.loads((tmp_path / result['path'] / 'candidates.json').read_text())
    assert saved['records'][0]['detail_refs'][0]['archive_path'] == relative
    assert '13800138000' not in saved['records'][0]['conclusion']
    assert '[REDACTED]' in saved['untrusted']['archive_path']
