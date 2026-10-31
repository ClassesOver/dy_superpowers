"""新版进度按当前模板问题、完整性及当前意图分流，各范围独立。"""
from math import ceil
from pathlib import Path
import pytest
from openpyxl import load_workbook
import server
from dy_superpowers_mcp.qa_templates import load_template
from dy_superpowers_mcp.qa_workbook import answer_column
from template_qa_helpers import save_answers


@pytest.fixture(autouse=True)
def isolated(tmp_path,monkeypatch):
    monkeypatch.setenv('CODEX_HOME',str(tmp_path/'codex'))


def call(root,tool='dy_research_progress',**args):
    result,error=server.execute_tool(tool,{'trace_id':'progress-regression','module':'expense',**args},root)
    assert error is None,error
    return result


@pytest.mark.parametrize('module',['expense','general-ledger'])
@pytest.mark.parametrize('complete',[True,False])
def test_complete_and_partial_confirmation_have_separate_progress(tmp_path,module,complete):
    save_answers(tmp_path,module,count=5,complete=complete)
    result=call(tmp_path,module=module,intent='progress')
    assert result['total']==len(load_template(module)['questions']) and result['confirmed_count']==5
    assert result['answered']==(5 if complete else 0)
    assert result['partial_count']==(0 if complete else 5)
    assert result['route']=='excel'


@pytest.mark.parametrize('module',['expense','general-ledger'])
@pytest.mark.parametrize('at_threshold',[False,True])
def test_70_percent_boundary_uses_full_completion(tmp_path,module,at_threshold):
    total=len(load_template(module)['questions'])
    count=ceil(total*0.7)-(not at_threshold)
    route='interactive' if at_threshold else 'excel'
    save_answers(tmp_path,module,count=count)
    result=call(tmp_path,module=module,intent='research')
    assert result['total']==total and result['answered']==count and result['route']==route
    assert result['next_action']==('ask_missing' if at_threshold else 'offer_checklist')


@pytest.mark.parametrize('has_answers',[False,True])
@pytest.mark.parametrize('intent,action',[('progress','report_progress'),('export','export_checklist'),('import','confirm_source'),('confirm','save_confirmed'),(None,'classify_intent')])
def test_mcp_intent_selects_action_without_overriding_threshold(tmp_path,has_answers,intent,action):
    count=ceil(len(load_template('expense')['questions'])*0.7) if has_answers else 0
    if count:save_answers(tmp_path,count=count)
    result=call(tmp_path,**({'intent':intent} if intent else {}))
    assert result['intent']==intent and result['next_action']==action
    assert result['route']==('interactive' if count else 'excel')
    event=next(e for e in server.TraceStore(tmp_path).read('progress-regression') if e['event']=='research.progress')
    assert event.get('intent')==(intent or 'unset') and event['next_action']==action


def test_partial_coverage_cannot_raise_complete_progress_above_threshold(tmp_path):
    count=ceil(len(load_template('expense')['questions'])*0.7)
    save_answers(tmp_path,count=count,complete=False)
    result=call(tmp_path,intent='research')
    assert result['confirmed_count']==count and result['answered']==0 and result['route']=='excel'


def test_unknown_intent_is_rejected(tmp_path):
    result,error=server.execute_tool('dy_research_progress',dict(trace_id='invalid',module='expense',intent='guess'),tmp_path)
    assert result is None and error['code']==-32602


@pytest.mark.parametrize('module',['expense','general-ledger'])
def test_empty_progress_is_read_only_and_uses_full_template(tmp_path,module):
    result=call(tmp_path,module=module)
    assert result['scope_mode']=='unscoped' and result['scope_id'] is None
    assert len(result['questions'])==result['total']==len(load_template(module)['questions'])
    assert not (tmp_path/'.dy_superpowers').exists()


def test_multiple_scopes_do_not_union_answers_and_export_separate_files(tmp_path):
    threshold=ceil(len(load_template('expense')['questions'])*0.7)
    first=save_answers(tmp_path,count=threshold-1,scope='甲院')
    second=save_answers(tmp_path,count=threshold,scope='乙院')
    result=call(tmp_path,intent='research')
    assert result['scope_mode']=='multiple' and result['answered']==0 and not result['records']
    assert sorted((s['answered'],s['route']) for s in result['scope_progress'])==[(threshold-1,'excel'),(threshold,'interactive')]
    exported=call(tmp_path,'dy_research_unanswered_export',confirmed=True)
    assert len(exported['files'])==2 and len({f['path'] for f in exported['files']})==2
    assert {f['scope_id'] for f in exported['files']}=={first['scope_id'],second['scope_id']}
    assert call(tmp_path,scope_id=first['scope_id'])['answered']==threshold-1


@pytest.mark.parametrize('count',[0,1])
def test_export_preserves_confirmation_and_full_question_set(tmp_path,count):
    if count:save_answers(tmp_path,count=count)
    _,error=server.execute_tool('dy_research_unanswered_export',dict(trace_id='denied',module='expense',confirmed=False),tmp_path)
    assert error['code']==-32602
    exported=call(tmp_path,'dy_research_unanswered_export',confirmed=True)
    template=load_template('expense')
    assert exported['question_count']==len(template['questions'])
    book=load_workbook(exported['path'])
    try:
        q=template['questions'][0]
        assert bool(book.active.cell(q['row'],answer_column(template)).value)==bool(count)
        assert all(book.active.cell(q['row'],answer_column(template)+1).value is None for q in template['questions'])
    finally:book.close()


@pytest.mark.parametrize('tool',['dy_research_progress','dy_research_unanswered_export'])
def test_corrupt_current_memory_cannot_be_hidden_by_unscoped_routing(tmp_path,tool):
    state=save_answers(tmp_path)
    path=Path(state['path']).with_name('memory.json')
    path.write_text(path.read_text().replace('本期按确认清单配置','未经确认编辑'))
    args=dict(trace_id='corrupt',module='expense')
    if tool.endswith('export'):args['confirmed']=True
    result,error=server.execute_tool(tool,args,tmp_path)
    assert result is None and error['code']==-32602


def test_explicit_unknown_scope_is_rejected_instead_of_invented(tmp_path):
    result,error=server.execute_tool('dy_research_unanswered_export',dict(trace_id='unknown',module='expense',scope_id='0123456789abcdef',confirmed=True),tmp_path)
    assert result is None and error['code']==-32602


def test_legacy_corrupt_records_do_not_enter_new_progress(tmp_path):
    path=tmp_path/'.dy_superpowers/modules/expense/research.md'
    path.parent.mkdir(parents=True)
    path.write_text('损坏的旧记录')
    assert call(tmp_path,intent='progress')['answered']==0
    assert path.read_text()=='损坏的旧记录'
