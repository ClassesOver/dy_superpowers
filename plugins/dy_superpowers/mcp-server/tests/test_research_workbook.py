"""当前答卷完整导出、候选回填、附件要求与并发下载边界。"""
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import pytest
from openpyxl import Workbook,load_workbook
from openpyxl.utils import get_column_letter
from dy_superpowers_mcp.qa_store import export_qa,parse_qa_workbook,read_qa
from dy_superpowers_mcp.qa_templates import load_template
from dy_superpowers_mcp.qa_workbook import answer_column
from template_qa_helpers import save_answers


@pytest.mark.parametrize('module',['expense','general-ledger'])
def test_initial_workbook_preserves_catalog_layout_and_attachment_headers(tmp_path,module):
    result=export_qa(tmp_path,module,confirmed=True)
    template=load_template(module)
    assert result['question_count']==len(template['questions'])
    book=load_workbook(result['path'])
    try:
        sheet=book.active
        expected_merges=set(template['layout']['merges'])
        for q in template['questions']:
            rows=q.get('rows',[q['row']])
            if len(rows)>1:
                for column in range(answer_column(template),answer_column(template)+3):
                    letter=get_column_letter(column)
                    expected_merges.add(f'{letter}{rows[0]}:{letter}{rows[-1]}')
        assert {str(r) for r in sheet.merged_cells.ranges}==expected_merges
        for q in template['questions']:
            assert sheet.cell(q['row'],4).value in {q['question'], None}
            confirmed_col=answer_column(template)
            assert sheet.cell(q['row'],confirmed_col).value is None
            assert sheet.cell(q['row'],confirmed_col+1).value is None
        for a in template['attachments']:
            assert [c.value for c in book[a['name']][1] if c.value is not None]==a['headers']
    finally:book.close()
    parsed=parse_qa_workbook(tmp_path,result['path'])
    assert not parsed['answers'] and parsed['can_write_confirmed'] is False


@pytest.mark.parametrize('module',['expense','general-ledger'])
def test_repeated_concurrent_downloads_do_not_overwrite(tmp_path,module):
    first=export_qa(tmp_path,module,confirmed=True)
    content=Path(first['path']).read_bytes()
    with ThreadPoolExecutor(max_workers=3) as pool:
        results=list(pool.map(lambda _:export_qa(tmp_path,module,confirmed=True),range(3)))
    assert len({r['path'] for r in results})==3
    assert Path(first['path']).read_bytes()==content


@pytest.mark.parametrize('module',['expense','general-ledger'])
def test_current_confirmed_cells_are_baselines_and_only_new_column_is_candidate(tmp_path,module):
    state=save_answers(tmp_path,module,count=2,complete=False)
    export=export_qa(tmp_path,module,confirmed=True)
    template=load_template(module);col=answer_column(template)+1
    book=load_workbook(export['path'])
    book.active.cell(template['questions'][1]['row'],col,'新确认候选')
    book.save(export['path']);book.close()
    parsed=parse_qa_workbook(tmp_path,export['path'])
    assert len(parsed['answers'])==1 and parsed['can_write_confirmed'] is False
    assert parsed['answers'][0]['question_id']==template['questions'][1]['question_id']
    assert parsed['answers'][0]['expected_revision']==state['records'][1]['revision']
    assert parsed['answers'][0]['research_complete'] is False
    assert read_qa(tmp_path,module)['answer_version']==state['answer_version']


def test_different_scopes_with_same_label_cannot_mix_downloads(tmp_path):
    one=save_answers(tmp_path,scope='甲院',answer='第一条\n第二条')
    two=save_answers(tmp_path,scope='乙院',answer='另一范围')
    result=export_qa(tmp_path,'expense',confirmed=True)
    assert len(result['files'])==2
    for item in result['files']:
        parsed=parse_qa_workbook(tmp_path,item['path'])
        assert parsed['scope_id'] in {one['scope_id'],two['scope_id']}
        assert not parsed['answers']


def test_attachment_requirements_are_exported_but_editing_is_only_candidate(tmp_path):
    state=save_answers(tmp_path)
    result=export_qa(tmp_path,'expense',confirmed=True)
    book=load_workbook(result['path'])
    sheet=book['已确认材料']
    question=state['questions'][0]
    assert sheet['A2'].value==question['config_item']
    assert sheet['B2'].value==question['question']
    sheet['D2']='新材料要求，待确认'
    book.save(result['path']);book.close()
    parsed=parse_qa_workbook(tmp_path,result['path'])
    assert parsed['attachment_changes'] and parsed['can_write_confirmed'] is False
    assert read_qa(tmp_path,'expense')['answer_version']==state['answer_version']


@pytest.mark.parametrize('change',['confirmed','metadata','question'])
def test_immutable_baseline_or_identity_edits_are_rejected(tmp_path,change):
    save_answers(tmp_path)
    result=export_qa(tmp_path,'expense',confirmed=True)
    book=load_workbook(result['path'])
    q=load_template('expense')['questions'][0]
    if change=='confirmed':book.active.cell(q['row'],answer_column(load_template('expense')),'伪造已确认')
    elif change=='question':book.active.cell(q['row'],4,'伪造题目')
    else:
        hidden=next(s for s in book if s.sheet_state=='veryHidden' or s.sheet_state=='hidden')
        hidden['A1']='伪造元数据'
    book.save(result['path']);book.close()
    with pytest.raises(ValueError):parse_qa_workbook(tmp_path,result['path'])


def test_old_layout_is_rejected_by_new_import_without_changing_memory(tmp_path):
    book=Workbook();book.active.append(['模块','检查顺序','调研子项','答案'])
    book.active.append(['expense','1','01','旧答案'])
    path=tmp_path/'旧格式.xlsx';book.save(path);book.close()
    with pytest.raises(ValueError):parse_qa_workbook(tmp_path,path)
    assert not (tmp_path/'.dy_superpowers').exists()


def test_export_requires_confirmation_and_refuses_memory_symlink(tmp_path):
    with pytest.raises(ValueError):export_qa(tmp_path,'expense',confirmed=False)
    other=tmp_path/'other';other.mkdir()
    (tmp_path/'.dy_superpowers').symlink_to(other,target_is_directory=True)
    with pytest.raises(ValueError):export_qa(tmp_path,'expense',confirmed=True)
