"""最终独立审查的三个实质边界；仅使用临时工作区。"""
import hashlib
import json
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

import server
from dy_superpowers_mcp import qa_workbook, report_pdf
from dy_superpowers_mcp.inspection import validate_inspection
from dy_superpowers_mcp.inspection_prepare import prepare_inspection
from dy_superpowers_mcp.qa_basis import optional_attachment
from dy_superpowers_mcp.qa_store import commit_qa, export_qa, read_qa
from dy_superpowers_mcp.qa_templates import load_template
from dy_superpowers_mcp.reports import export_report_pdf, read_report, save_report
from template_qa_helpers import make_report


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setenv('CODEX_HOME', str(tmp_path / 'codex'))
    monkeypatch.delenv('DB_NAME', raising=False)
    monkeypatch.delenv('DY_ODOO_ENVIRONMENT', raising=False)


def _manual_report(root, qid, answer):
    template = load_template('general-ledger')
    state = commit_qa(root, 'general-ledger', scope='甲院/一期', template_revision=template['revision'],
        answers=[dict(question_id=qid, answer=answer, research_complete=True, expected_revision=None, attachments=[])],
        batch_id='人工签认', confirmed=True, confirmed_by='测试')
    report = make_report(root, 'general-ledger', state=state)
    signed = root / '签认.txt'
    signed.write_text('本期该功能及对应附件不适用，已由业务负责人签认', encoding='utf-8')
    evidence = dict(kind='manual', reference=signed.name, sha256=hashlib.sha256(signed.read_bytes()).hexdigest(),
        signed_by='业务负责人', signed_at=datetime.now(timezone.utc).isoformat())
    record = next(r for r in report['records'] if r['question_id'] == qid)
    record.update(result='不适用', verified_at=datetime.now(timezone.utc).isoformat(), evidence_refs=[evidence],
        requirements=[dict(requirement_id='不适用范围', text=answer, result='不适用', verification_target='document',
            source_refs=[dict(kind='answer', question_id=qid, revision=record['answer_revision'])], evidence_refs=[evidence])])
    return state, report


def test_optional_attachment_explicitly_declined_with_manual_proof(tmp_path):
    state, report = _manual_report(tmp_path, 'GL-2-010', '本期不限制用户使用特定科目，科目授权清单不适用')
    prepared = prepare_inspection(tmp_path, 'general-ledger', state['scope_id'], 'UAT')
    assert next(q for q in prepared['items'] if q['question_id'] == 'GL-2-010')['readiness'] == 'research_ready'
    validated = validate_inspection(tmp_path, report)
    assert validated['valid'], validated['errors']


def test_current_template_attachment_instruction_and_question_are_runtime_basis(tmp_path):
    assert optional_attachment({'attachment_instruction': '附件（按需）：预算科目类型清单'})
    template = load_template('expense')
    state = commit_qa(tmp_path, 'expense', scope='甲院', template_revision=template['revision'],
        answers=[dict(question_id=template['questions'][1]['question_id'], answer='局部事实',
                      research_complete=False, expected_revision=None)],
        batch_id='当前模板', confirmed=True)
    prepared = prepare_inspection(tmp_path, 'expense', state['scope_id'], 'UAT')
    assert prepared['items'][0]['next_question'] == template['questions'][0]['question']


def test_optional_attachment_without_real_manual_evidence_still_rejected(tmp_path):
    _, report = _manual_report(tmp_path, 'GL-2-010', '本期科目授权清单不适用')
    record = next(r for r in report['records'] if r['question_id'] == 'GL-2-010')
    record['evidence_refs'] = []
    record['requirements'][0]['evidence_refs'] = []
    assert not validate_inspection(tmp_path, report)['valid']


@pytest.mark.parametrize('qid,answer', [
    ('GL-2-010', '本期需要科目授权，应按清单配置'),
    ('GL-1-001', '本期不使用总账基础及凭证应用，无需附件'),
    ('GL-2-006', '本期财务和预算科目类型均不适用'),
])
def test_required_attachment_or_applicable_branch_still_blocks(tmp_path, qid, answer):
    state, report = _manual_report(tmp_path, qid, answer)
    prepared = prepare_inspection(tmp_path, 'general-ledger', state['scope_id'], 'UAT')
    item = next(q for q in prepared['items'] if q['question_id'] == qid)
    if qid in ('GL-1-001', 'GL-2-006'):
        assert item['readiness'] == 'attachment_missing'
    else:
        record = next(r for r in report['records'] if r['question_id'] == qid)
        record['result'] = '通过'
        record['requirements'][0]['result'] = '通过'
    assert not validate_inspection(tmp_path, report)['valid']


@pytest.mark.parametrize('value,expected', [
    (datetime(2026, 10, 1), '2026-10-01T00:00:00'),
    (date(2026, 10, 1), '2026-10-01T00:00:00'),
    (time(9, 30, 15), '09:30:15'),
    (timedelta(hours=36), 'P1DT12H0M0S'),
    (2026, 2026), (1.25, 1.25), (True, True),
])
def test_attachment_excel_scalars_survive_mcp_serialization_readonly(tmp_path, value, expected):
    exported = export_qa(tmp_path, 'general-ledger', confirmed=True)
    path = Path(exported['path'])
    book = load_workbook(path)
    book['会计年度与分配公司清单']['D2'] = value
    if isinstance(value, timedelta):
        book['会计年度与分配公司清单']['D2'].number_format = '[hh]:mm:ss'
    book.save(path)
    before = {str(p.relative_to(tmp_path)): p.read_bytes() for p in tmp_path.rglob('*') if p.is_file()}
    response = server.handle(dict(jsonrpc='2.0', id=1, method='tools/call', params=dict(
        name='dy_research_answers_import', arguments=dict(trace_id='date-repro', path=str(path)))), tmp_path)
    json.dumps(response)
    assert 'error' not in response, response
    result = json.loads(response['result']['content'][0]['text'])
    candidate = next(c for c in result['attachment_changes'] if c['cell'] == 'D2')
    assert candidate['sheet'] == '会计年度与分配公司清单'
    assert candidate['value'] == expected and candidate['before'] is None
    assert not candidate['can_write_confirmed'] and not result['can_write_confirmed']
    assert result['answers'] == [] and read_qa(tmp_path, 'general-ledger')['records'] == []
    for relative, content in before.items():
        assert (tmp_path / relative).read_bytes() == content
    traces = ''.join(p.read_text() for p in tmp_path.rglob('*.jsonl'))
    assert 'attachment_changes' not in traces and '年度开始日期' not in traces


def test_workbook_semantic_baseline_dates_are_json_stable():
    book = Workbook()
    book.active.append([date(2026, 10, 1), datetime(2026, 10, 1, 12, 30), time(9, 30)])
    baseline = qa_workbook.describe(book)
    assert baseline == json.loads(json.dumps(baseline))
    assert baseline[0]['rows'][0] == ['2026-10-01', '2026-10-01T12:30:00', '09:30:00']


def _advance_answer(root, report):
    state = read_qa(root, 'expense', report['basis']['scope_id'])
    answer = state['records'][0]
    changed = commit_qa(root, 'expense', scope=state['scope'], scope_id=state['scope_id'],
        template_revision=state['template_revision'], batch_id='渲染期间确认新答案', confirmed=True,
        answers=[dict(question_id=answer['question_id'], answer='渲染期间已确认新业务要求', research_complete=True,
                      expected_revision=answer['revision'])])
    assert changed['answer_version'] != report['basis']['answer_version']


@pytest.mark.parametrize('entry', ['mcp', 'direct'])
def test_report_rechecks_basis_after_pdf_render_before_publication(tmp_path, monkeypatch, entry):
    report = make_report(tmp_path)
    validation = validate_inspection(tmp_path, report)
    assert validation['valid']
    original = report_pdf.render_report_pdf
    def render_and_confirm(*args, **kwargs):
        content = original(*args, **kwargs)
        _advance_answer(tmp_path, report)
        return content
    monkeypatch.setattr(report_pdf, 'render_report_pdf', render_and_confirm)
    if entry == 'mcp':
        result, error = server.execute_tool('dy_inspection_report_save', dict(trace_id='race', report=report, confirmed=True), tmp_path)
        assert error is not None and not (result and result.get('saved'))
    else:
        with pytest.raises(ValueError, match='版本|依据|校验'):
            save_report(tmp_path, 'expense', report, validation)
    assert not list((tmp_path / '.dy_superpowers' / 'reports').rglob('*.md'))
    assert not list((tmp_path / '.dy_superpowers' / 'reports').rglob('*.pdf'))


def test_pdf_reexport_rechecks_basis_and_keeps_historical_markdown(tmp_path, monkeypatch):
    report = make_report(tmp_path)
    report_id, md_path, pdf_path = save_report(tmp_path, 'expense', report, validate_inspection(tmp_path, report))
    historical = md_path.read_bytes()
    pdf_path.unlink()
    original = report_pdf.render_report_pdf
    def render_and_confirm(*args, **kwargs):
        content = original(*args, **kwargs)
        _advance_answer(tmp_path, report)
        return content
    monkeypatch.setattr(report_pdf, 'render_report_pdf', render_and_confirm)
    with pytest.raises(ValueError, match='版本|依据|校验'):
        export_report_pdf(tmp_path, 'expense', report_id)
    assert not pdf_path.exists()
    assert md_path.read_bytes() == historical
    assert read_report(tmp_path, 'expense', report_id) == historical.decode().split('\n## 机器可读数据\n', 1)[0].rstrip() + '\n'
