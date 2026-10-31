"""模板目录必须独立于项目样例，并完整保留用户指定表的业务内容。"""

import hashlib
import importlib
import importlib.util
import json
import re
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[4]
PACKAGE = ROOT / "plugins/dy_superpowers/mcp-server/src/dy_superpowers_mcp"
REFERENCES = ROOT / "plugins/dy_superpowers/skills/module-research/references"
GENERATOR = ROOT / "scripts/generate-research-templates.py"
SOURCES = {
    "expense": ("费用模块调研清单.xlsx", 2, 5, 25, 25),
    "general-ledger": ("总账模块调研模板.xlsx", 2, 5, 37, 39),
}


def catalog(module):
    """缺少接口时以功能断言失败，避免导入错误掩盖 RED 原因。"""
    assert importlib.util.find_spec("dy_superpowers_mcp.qa_templates"), "缺少模板目录接口"
    return importlib.import_module("dy_superpowers_mcp.qa_templates").load_template(module)


def source_book(module):
    path = REFERENCES / SOURCES[module][0]
    if not path.exists():
        pytest.skip("用户源工作簿未随插件发布")
    return load_workbook(path, data_only=False)


def expanded_value(sheet, row, column):
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= column <= merged.max_col:
            return sheet.cell(merged.min_row, merged.min_col).value
    return sheet.cell(row, column).value


@pytest.mark.parametrize("module", SOURCES)
def test_catalog_has_every_unique_point_and_original_group(module):
    template = catalog(module)
    _, header_row, _, group_count, attachment_count = SOURCES[module]
    assert template["module"] == module
    assert len(template["questions"]) == group_count
    assert len({point["question_id"] for point in template["questions"]}) == group_count
    assert all(point['question_id'] == point['check_order'] for point in template['questions'])
    assert len({point["check_order"] for point in template["questions"]}) == group_count
    assert len(template["attachments"]) == attachment_count
    assert template["layout"]["header_row"] == header_row
    assert set(template) == {"module", "revision", "title", "questions", "attachments", "layout"}


@pytest.mark.parametrize("module", SOURCES)
def test_every_question_matches_source_text_and_effective_attachment_link(module):
    book = source_book(module)
    sheet = book.worksheets[0]
    template = catalog(module)
    grouped_rows = {}
    for row in range(SOURCES[module][1] + 1, sheet.max_row + 1):
        grouped_rows.setdefault(expanded_value(sheet, row, 1), []).append(row)
    for ordinal, (point, (check, rows)) in enumerate(zip(template['questions'], grouped_rows.items(), strict=True), 1):
        row = rows[0]
        instructions, names = [], []
        for source_row in rows:
            cell = sheet.cell(source_row, 5)
            if cell.value and cell.value not in instructions:
                instructions.append(cell.value)
            link = cell.hyperlink
            target = (link.location or link.target) if link else None
            name = target.lstrip('#').rsplit('!', 1)[0].strip("'") if target else None
            if name in book.sheetnames[1:] and name not in names:
                names.append(name)
        assert point == {
            "question_id": check,
            "check_order": check,
            "research_id": f"{ordinal:02d}",
            "ordinal": ordinal,
            "row": row,
            "rows": rows,
            "config_item": expanded_value(sheet, row, 2),
            "config_path": expanded_value(sheet, row, 3),
            "question": expanded_value(sheet, row, 4),
            "extraction_guide": "",
            "reference_answer": "",
            "attachment_instruction": '\n'.join(instructions),
            "attachment_names": names,
            "reuse_question_id": {'GL-2-009': 'GL-2-007', 'GL-3-012': 'GL-3-011'}.get(check),
        }


@pytest.mark.parametrize("module", SOURCES)
def test_attachment_headers_match_all_real_columns_without_sample_rows(module):
    book = source_book(module)
    attachments = catalog(module)["attachments"]
    assert [entry["name"] for entry in attachments] == book.sheetnames[1:]
    for attachment, sheet in zip(attachments, book.worksheets[1:], strict=True):
        headers = [cell.value for cell in sheet[2 if module == "expense" else 1] if cell.value is not None]
        assert attachment == {"name": sheet.title, "headers": headers}


@pytest.mark.parametrize("module", SOURCES)
def test_layout_preserves_main_structure_and_excludes_spare_answers(module):
    sheet = source_book(module).worksheets[0]
    layout = catalog(module)["layout"]
    _, header_row, business_columns, _, _ = SOURCES[module]
    columns = 9 if module == "expense" else 6
    assert layout["headers"] == [sheet.cell(header_row, column).value for column in range(1, columns + 1)]
    assert layout["rows"] == [
        [sheet.cell(row, column).value if column <= business_columns else None for column in range(1, columns + 1)]
        for row in range(1, sheet.max_row + 1)
    ]
    assert set(layout["merges"]) == {str(merged) for merged in sheet.merged_cells.ranges}
    for column in range(1, columns + 1):
        letter = get_column_letter(column)
        dimension = next(
            dimension for dimension in sheet.column_dimensions.values()
            if dimension.min <= column <= dimension.max
        )
        assert layout["widths"][letter] == dimension.width


def test_attachment_branches_are_one_business_question_without_duplicate_answers():
    points = [point for point in catalog("general-ledger")["questions"] if point["check_order"] == "GL-8-024"]
    assert len(points) == 1
    assert points[0]['question_id'] == 'GL-8-024'
    assert points[0]['rows'] == [48, 49]
    assert points[0]['attachment_instruction'].count('01｜') == 2


def test_attachment_alias_uses_real_sheet_and_does_not_invent_template():
    template = catalog("general-ledger")
    point = next(point for point in template["questions"] if point["question_id"] == "GL-6-021")
    assert "预算凭证模板清单" in point["attachment_instruction"]
    assert point["reference_answer"] == ""
    assert point["attachment_names"] == ["凭证模板清单"]
    assert "预算凭证模板清单" not in {entry["name"] for entry in template["attachments"]}


def test_load_returns_independent_nested_values_and_rejects_unknown_module():
    first = catalog("expense")
    first["questions"][0]["attachment_names"].clear()
    first["attachments"][0]["headers"][0] = "污染"
    first["layout"]["rows"][0][0] = "污染"
    second = catalog("expense")
    assert second["questions"][0]["attachment_names"] == ["01_应用安装"]
    assert second["attachments"][0]["headers"][0] == "序号"
    assert second["layout"]["rows"][0][0] == "费用模块调研清单"
    with pytest.raises(ValueError):
        catalog("unknown")


@pytest.mark.parametrize("module", SOURCES)
def test_revision_is_normalized_content_sha256(module):
    template = catalog(module)
    on_disk = json.loads((PACKAGE / f"data/research_templates/{module}.json").read_text())
    assert on_disk == template
    revision = on_disk.pop("revision")
    canonical = json.dumps(on_disk, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    assert revision == hashlib.sha256(canonical.encode()).hexdigest()


@pytest.mark.parametrize("module", SOURCES)
def test_question_projection_keeps_every_question_without_answers_or_attachment_schema(module):
    template = catalog(module)
    guide = (REFERENCES / f"{module}-research-question.md").read_text()
    assert re.findall(r"^## (.+)$", guide, re.MULTILINE) == ["模块 QA"]
    groups = {point["check_order"]: point for point in template["questions"]}
    assert re.findall(r"^### (.+)$", guide, re.MULTILINE) == [
        f"{point['check_order']} {point['config_item']}" for point in groups.values()
    ]
    expected_lines = [f"# {template['title']} QA 指引", "", "## 模块 QA", ""]
    for point in groups.values():
        expected_lines.extend([f"### {point['check_order']} {point['config_item']}", "",
                               point['question'].replace("\n", "<br>"), ""])
    assert guide == "\n".join(expected_lines)
    assert not re.findall(r"^####", guide, re.MULTILINE)
    for label in ("提问要素：", "配置路径：", "结论参考：", "参考答案：", "附件：", "## 附件字段", "## 附件清单"):
        assert not any(line.startswith(label) for line in guide.splitlines())



def test_generator_reproduces_catalog_projection_and_audit_without_copying_workbooks(tmp_path):
    for module in SOURCES:
        source_book(module)
    assert GENERATOR.exists(), "缺少可重复生成脚本"
    result = subprocess.run(
        [sys.executable, str(GENERATOR), "--source-dir", str(REFERENCES), "--output-root", str(tmp_path)],
        capture_output=True, text=True,
    )
    assert result.returncode == 0, result.stderr
    for module in SOURCES:
        for relative in (
            f"plugins/dy_superpowers/mcp-server/src/dy_superpowers_mcp/data/research_templates/{module}.json",
            f"plugins/dy_superpowers/skills/module-research/references/{module}-research-question.md",
        ):
            assert (tmp_path / relative).read_bytes() == (ROOT / relative).read_bytes()
    audit = "docs/research-sources/template-qa-20261030.md"
    assert (tmp_path / audit).read_bytes() == (ROOT / audit).read_bytes()
    assert not list(tmp_path.rglob("*.xlsx"))


@pytest.mark.parametrize("module", SOURCES)
def test_generator_never_promotes_answer_or_spare_columns_to_questions_or_references(tmp_path, module):
    spec = importlib.util.spec_from_file_location('research_template_generator', GENERATOR)
    generator = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(generator)
    book = source_book(module)
    sheet = book.worksheets[0]
    sheet['F3'] = '不得读取的项目答案'
    if module == 'expense':
        sheet['G3'] = '不得读取的旧提问'
        sheet['H3'] = '不得读取的旧参考'
    path = tmp_path / '回填副本.xlsx'
    book.save(path)
    template, _ = generator.extract_template(module, path)
    assert len(template['questions']) == SOURCES[module][3]
    assert all(q['extraction_guide'] == q['reference_answer'] == '' for q in template['questions'])
    assert all(all(value is None for value in row[5:]) for row in template['layout']['rows'])
    assert '不得读取' not in json.dumps(template, ensure_ascii=False)
    guide = generator.render_guide(template)
    assert '不得读取' not in guide
    assert all(q['attachment_instruction'] not in guide for q in template['questions'] if q['attachment_instruction'])
