"""答卷回答区遵循发布模板的答案表头，不能按模块硬编码。"""
from copy import deepcopy
import pytest
from openpyxl import load_workbook
from dy_superpowers_mcp.qa_templates import load_template
from dy_superpowers_mcp.qa_workbook import answer_column, build_workbook


@pytest.mark.parametrize('module',['expense','general-ledger'])
def test_answer_column_tracks_current_layout_answer_header(tmp_path,module):
    template=deepcopy(load_template(module))
    template['layout']['headers']=['检查编号','配置项','配置路径','问题','附件清单','答案']
    assert answer_column(template)==6
    path=tmp_path/'答卷.xlsx'
    build_workbook(path,template,[],{})
    book=load_workbook(path)
    try:
        header=template['layout']['header_row']
        assert [book.active.cell(header,col).value for col in (6,7,8)]==['已确认答案','本轮补充或修订','调研状态']
    finally:book.close()
