"""共用 QA 的确认、版本、工作簿与附件边界。"""

import hashlib
import importlib
import importlib.util
import json
import multiprocessing
from pathlib import Path

import pytest
from openpyxl import load_workbook

from dy_superpowers_mcp.qa_templates import load_template
from dy_superpowers_mcp.qa_workbook import answer_column


def api():
    assert importlib.util.find_spec('dy_superpowers_mcp.qa_store'), '缺少版本化 QA 存储'
    return importlib.import_module('dy_superpowers_mcp.qa_store')


def item(module='expense', index=0, answer='本期交付移动报销', revision=None, **extra):
    return {'question_id': load_template(module)['questions'][index]['question_id'],
            'answer': answer, 'research_complete': True, 'expected_revision': revision, **extra}


def save(workspace, answers, *, module='expense', scope='甲院/本期', scope_id=None, batch='批次一', **extra):
    return api().commit_qa(workspace, module, scope=scope, scope_id=scope_id,
                           template_revision=load_template(module)['revision'], answers=answers,
                           batch_id=batch, confirmed=True, **extra)


def records(result):
    return {record['question_id']: record for record in result['records']}


@pytest.mark.parametrize('module', ['expense', 'general-ledger'])
def test_empty_read_is_pure_and_never_counts_template_examples(tmp_path, module):
    result = api().read_qa(tmp_path, module)
    total = len(load_template(module)['questions'])
    assert result['total'] == total and result['answered'] == result['confirmed_count'] == 0
    assert result['route'] == 'excel' and result['records'] == []
    assert len(result['questions']) == total and result['answer_version'] is None
    assert list(tmp_path.iterdir()) == []


def test_partial_progress_and_seventy_percent_boundary(tmp_path):
    rows = [item(index=i) for i in range(17)]
    first = save(tmp_path, rows)
    assert first['answered'] == 17 and first['route'] == 'excel'
    second = save(tmp_path, [item(index=17)], batch='补充')
    assert second['answered'] == 18 and second['route'] == 'interactive'
    partial = save(tmp_path, [item(index=18, research_complete=False)], batch='局部')
    assert partial['answered'] == 18 and partial['confirmed_count'] == 19
    assert partial['partial_count'] == 1 and '完整完成 18/25' in partial['progress_summary']


def test_replacement_blank_preservation_and_batch_idempotency(tmp_path):
    first = save(tmp_path, [item(), item(index=1, answer='启用OCR')])
    q = item()['question_id']
    revision = records(first)[q]['revision']
    replay = save(tmp_path, [item(), item(index=1, answer='启用OCR')])
    assert replay['answer_version'] == first['answer_version']
    with pytest.raises(ValueError, match='批次'):
        save(tmp_path, [item(answer='不同正文')])
    changed = save(tmp_path, [item(answer='本期不交付移动报销', revision=revision)], batch='替换')
    assert records(changed)[q]['answer'] == '本期不交付移动报销'
    assert len(changed['records']) == 2
    current = records(changed)[q]['revision']
    unchanged = save(tmp_path, [item(answer='  ', revision=current)], batch='空白')
    assert unchanged['answer_version'] == changed['answer_version']
    same = save(tmp_path, [item(answer='本期不交付移动报销', revision=current)], batch='相同')
    assert same['answer_version'] == changed['answer_version']
    assert save(tmp_path, [], batch='空批')['answer_version'] == changed['answer_version']


def test_conflict_rejects_whole_batch_but_other_question_can_merge(tmp_path):
    first = save(tmp_path, [item()])
    old = records(first)[item()['question_id']]['revision']
    save(tmp_path, [item(answer='已更新', revision=old)], batch='聊天更改')
    with pytest.raises(ValueError, match='冲突'):
        save(tmp_path, [item(answer='旧版本覆盖', revision=old), item(index=1)], batch='冲突整批')
    assert len(api().read_qa(tmp_path, 'expense')['records']) == 1
    assert len(save(tmp_path, [item(index=1)], batch='另一题')['records']) == 2


@pytest.mark.parametrize('module', ['expense','general-ledger'])
def test_chat_download_fill_commit_download_preserves_template(tmp_path, module):
    confirmed_col = answer_column(load_template(module))
    new_col, status_col = confirmed_col + 1, confirmed_col + 2
    first = save(tmp_path, [item(module=module)], module=module)
    exported = api().export_qa(tmp_path, module, confirmed=True)
    book = load_workbook(exported['path'])
    sheet = book.worksheets[0]
    template = load_template(module)
    row = template['questions'][0]['row']
    assert sheet.cell(row, confirmed_col).value == '本期交付移动报销'
    assert sheet.cell(row, new_col).value is None
    assert sheet.cell(row, status_col).value == '已完成'
    assert sheet.cell(row + 1, confirmed_col).value is None
    assert set(template['layout']['merges']) <= set(map(str, sheet.merged_cells.ranges))
    assert sheet.max_row == len(template['layout']['rows'])
    for attachment in template['attachments']:
        assert [cell.value for cell in book[attachment['name']][1]][:len(attachment['headers'])] == attachment['headers']
        assert book[attachment['name']].max_row == 1
    next_row = template['questions'][1]['row']
    sheet.cell(next_row, new_col, '本期已确认启用OCR')
    book.save(exported['path'])
    parsed = api().parse_qa_workbook(tmp_path, exported['path'])
    assert parsed['can_write_confirmed'] is False and len(parsed['answers']) == 1
    assert parsed['base_version'] == first['answer_version']
    assert parsed['answers'][0]['expected_revision'] is None
    assert parsed['answers'][0]['research_complete'] is False
    confirmed_answers = [{**answer, 'research_complete': True} for answer in parsed['answers']]
    saved = save(tmp_path, confirmed_answers, module=module, scope_id=parsed['scope_id'], batch='回填')
    again = api().export_qa(tmp_path, module, confirmed=True)
    assert again['path'] != exported['path']
    assert load_workbook(again['path']).worksheets[0].cell(next_row, confirmed_col).value == '本期已确认启用OCR'
    assert len(saved['records']) == 2


def test_old_export_uses_per_question_revision_and_can_continue(tmp_path):
    initial = save(tmp_path, [item()])
    exported = api().export_qa(tmp_path, 'expense', confirmed=True)
    old = records(initial)[item()['question_id']]['revision']
    save(tmp_path, [item(answer='聊天修订', revision=old)], batch='聊天')
    book = load_workbook(exported['path'])
    book.worksheets[0].cell(load_template('expense')['questions'][1]['row'], answer_column(load_template('expense')) + 1, '另一题答卷')
    book.save(exported['path'])
    parsed = api().parse_qa_workbook(tmp_path, exported['path'])
    confirmed_answers = [{**answer, 'research_complete': True} for answer in parsed['answers']]
    assert save(tmp_path, confirmed_answers, batch='续答')['answered'] == 2
    book.worksheets[0].cell(load_template('expense')['questions'][0]['row'], answer_column(load_template('expense')) + 1, '旧答卷同题')
    book.save(exported['path'])
    parsed = api().parse_qa_workbook(tmp_path, exported['path'])
    with pytest.raises(ValueError, match='冲突'):
        save(tmp_path, parsed['answers'], batch='同题冲突')


@pytest.mark.parametrize('mutation', ['question', 'confirmed', 'metadata', 'status'])
def test_parse_rejects_forged_or_modified_immutable_cells(tmp_path, mutation):
    save(tmp_path, [item()])
    exported = api().export_qa(tmp_path, 'expense', confirmed=True)
    book = load_workbook(exported['path'])
    if mutation == 'metadata':
        meta = book['_dy_qa']
        data = json.loads(meta['A1'].value)
        data['base_version'] = '伪造版本'
        meta['A1'] = json.dumps(data)
    else:
        col = {'question': 4, 'confirmed': answer_column(load_template('expense')), 'status': answer_column(load_template('expense')) + 2}[mutation]
        book.worksheets[0].cell(3, col, '擅自修改')
    book.save(exported['path'])
    with pytest.raises(ValueError):
        api().parse_qa_workbook(tmp_path, exported['path'])


def test_unregistered_legacy_and_cross_workspace_workbooks_are_rejected(tmp_path):
    from openpyxl import Workbook
    legacy = tmp_path / '旧答卷.xlsx'
    Workbook().save(legacy)
    with pytest.raises(ValueError):
        api().parse_qa_workbook(tmp_path, legacy)
    exported = api().export_qa(tmp_path, 'expense', confirmed=True)
    other = tmp_path / '其他工作区'
    other.mkdir()
    copied = other / '复制.xlsx'
    copied.write_bytes(Path(exported['path']).read_bytes())
    with pytest.raises(ValueError):
        api().parse_qa_workbook(other, copied)


def test_scope_and_module_are_isolated_and_multiple_exports_are_separate(tmp_path):
    one = save(tmp_path, [item(answer='甲院答案')])
    save(tmp_path, [item(answer='乙院答案')], scope='乙院', batch='乙院')
    save(tmp_path, [item(module='general-ledger')], module='general-ledger', batch='总账')
    mixed = api().read_qa(tmp_path, 'expense')
    assert mixed['scope_mode'] == 'multiple' and mixed['records'] == []
    assert len(mixed['scope_progress']) == 2
    assert records(api().read_qa(tmp_path, 'expense', one['scope_id']))[item()['question_id']]['answer'] == '甲院答案'
    assert len(api().export_qa(tmp_path, 'expense', confirmed=True)['files']) == 2
    with pytest.raises(ValueError, match='范围'):
        save(tmp_path, [item(index=1)], scope='未知范围', scope_id=one['scope_id'], batch='错范围')
    with pytest.raises(ValueError):
        api().read_qa(tmp_path, 'expense', '../其他')


def attachment(workspace, text='只启用甲院移动报销'):
    path = workspace / '材料.txt'
    path.write_text(text)
    return {'reference': '材料.txt', 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
            'requirements': ['甲院启用移动报销']}


def test_attachment_digest_revision_preservation_and_drift(tmp_path):
    bound = attachment(tmp_path)
    first = save(tmp_path, [item(attachments=[bound])])
    q = item()['question_id']
    record = records(first)[q]
    changed = save(tmp_path, [item(answer='文字修改', revision=record['revision'])], batch='保留附件')
    assert records(changed)[q]['attachments'] == [bound]
    attachment(tmp_path, '已修订材料')
    read = api().read_qa(tmp_path, 'expense')
    assert read['attachment_integrity'] is False
    assert read['attachment_changes'][0]['question_id'] == q
    bound['sha256'] = hashlib.sha256((tmp_path / '材料.txt').read_bytes()).hexdigest()
    updated = save(tmp_path, [item(answer='文字修改', revision=records(changed)[q]['revision'], attachments=[bound])], batch='确认新附件')
    assert records(updated)[q]['revision'] != records(changed)[q]['revision']
    cleared = save(tmp_path, [item(answer='文字修改', revision=records(updated)[q]['revision'], attachments=[])], batch='确认清除')
    assert records(cleared)[q]['attachments'] == []


def test_attachment_export_and_changes_are_candidates_only(tmp_path):
    bound = attachment(tmp_path)
    save(tmp_path, [item(attachments=[bound])])
    exported = api().export_qa(tmp_path, 'expense', confirmed=True)
    book = load_workbook(exported['path'])
    sheet = book['01_应用安装']
    assert '甲院启用移动报销' in [cell.value for cell in sheet[2]]
    sheet['B2'] = '申请新增OCR应用'
    book.save(exported['path'])
    parsed = api().parse_qa_workbook(tmp_path, exported['path'])
    assert parsed['answers'] == [] and parsed['attachment_changes']
    change = parsed['attachment_changes'][0]
    assert change['can_write_confirmed'] is False and change['sheet'] == '01_应用安装'
    assert change['scope_id'] == parsed['scope_id']
    assert api().read_qa(tmp_path, 'expense')['records'][0]['attachments'] == [bound]


@pytest.mark.parametrize('bad', ['password=abcdef', '手机号：13812345678', 'email=person@example.com', '张三身份证号110101199001011234'])
def test_sensitive_answers_are_rejected_without_writing(tmp_path, bad):
    with pytest.raises(ValueError, match='敏感'):
        save(tmp_path, [item(answer=bad)])
    assert list(tmp_path.iterdir()) == []


def test_attachment_validation_and_confirmation_are_strict(tmp_path):
    bound = attachment(tmp_path)
    with pytest.raises(ValueError):
        save(tmp_path, [item(attachments=[{**bound, 'sha256': '0' * 64}])])
    with pytest.raises(ValueError):
        save(tmp_path, [item(attachments=[{**bound, 'reference': '../材料.txt'}])])
    with pytest.raises(ValueError, match='敏感'):
        save(tmp_path, [item(attachments=[{**bound, 'requirements': ['token=secret-value']}])])
    with pytest.raises(ValueError, match='确认'):
        api().export_qa(tmp_path, 'expense', confirmed=False)
    with pytest.raises(ValueError, match='确认'):
        api().commit_qa(tmp_path, 'expense', scope='甲院', template_revision=load_template('expense')['revision'],
                        answers=[item()], batch_id='未确认', confirmed=False)


def test_symlinked_storage_and_attachment_paths_are_rejected(tmp_path):
    outside = tmp_path.parent / (tmp_path.name + '-outside')
    outside.mkdir()
    (tmp_path / '.dy_superpowers').symlink_to(outside, target_is_directory=True)
    with pytest.raises(ValueError):
        save(tmp_path, [item()])
    assert list(outside.iterdir()) == []


@pytest.mark.parametrize('failure', ['generate', 'readback'])
def test_failed_snapshot_does_not_publish_pointer(tmp_path, monkeypatch, failure):
    first = save(tmp_path, [item()])
    workbook = importlib.import_module('dy_superpowers_mcp.qa_workbook')
    def fail(*args, **kwargs):
        raise ValueError('注入工作簿失败')
    monkeypatch.setattr(workbook, 'build_workbook' if failure == 'generate' else 'verify_snapshot', fail)
    with pytest.raises(ValueError, match='注入'):
        save(tmp_path, [item(index=1)], batch='失败')
    monkeypatch.undo()
    assert api().read_qa(tmp_path, 'expense')['answer_version'] == first['answer_version']


@pytest.mark.parametrize('filename', ['memory.json', 'memory.md', 'answers.xlsx', 'manifest.json'])
def test_corrupt_current_snapshot_is_rejected_without_fallback(tmp_path, filename):
    save(tmp_path, [item()])
    result = save(tmp_path, [item(index=1)], batch='第二版')
    current = Path(result['path']).parent
    path = current / filename
    path.write_bytes(path.read_bytes() + '\n损坏'.encode())
    with pytest.raises(ValueError):
        api().read_qa(tmp_path, 'expense')


def _worker(workspace, index, queue):
    try:
        queue.put(('ok', save(Path(workspace), [item(index=index)], batch=f'进程{index}')['answer_version']))
    except Exception as exc:
        queue.put(('error', repr(exc)))


def test_two_process_publications_merge_different_questions(tmp_path):
    api()
    context = multiprocessing.get_context('fork')
    queue = context.Queue()
    children = [context.Process(target=_worker, args=(str(tmp_path), index, queue)) for index in (0, 1)]
    for child in children:
        child.start()
    for child in children:
        child.join(20)
        assert child.exitcode == 0
    results = [queue.get(timeout=2) for _ in children]
    assert all(status == 'ok' for status, _ in results), results
    result = api().read_qa(tmp_path, 'expense')
    assert result['answered'] == 2 and len(result['records']) == 2


def test_version_workbook_copy_is_registered_and_parsing_has_no_writes(tmp_path):
    saved = save(tmp_path, [item()])
    copied = tmp_path / '重新命名答卷.xlsx'
    copied.write_bytes(Path(saved['path']).read_bytes())
    book = load_workbook(copied)
    book.worksheets[0].cell(load_template('expense')['questions'][1]['row'], answer_column(load_template('expense')) + 1, '新增第二问题')
    book.save(copied)
    before = {str(p.relative_to(tmp_path)): p.read_bytes() for p in tmp_path.rglob('*') if p.is_file()}
    parsed = api().parse_qa_workbook(tmp_path, copied)
    after = {str(p.relative_to(tmp_path)): p.read_bytes() for p in tmp_path.rglob('*') if p.is_file()}
    assert before == after
    assert parsed['answers'] == [item(index=1, answer='新增第二问题', research_complete=False)]
    assert parsed['scope'] == '甲院/本期'


def test_attachment_requirements_use_explicit_fields_instead_of_fabricated_business_cells(tmp_path):
    bound = attachment(tmp_path)
    save(tmp_path, [item(attachments=[bound])])
    exported = api().export_qa(tmp_path, 'expense', confirmed=True)
    sheet = load_workbook(exported['path'])['01_应用安装']
    headers = [cell.value for cell in sheet[1]]
    assert headers[-4:] == ['配置项', '调研问题', '材料名称', '已确认要求']
    assert sheet.cell(2, 1).value is None  # 原字段“序号”不能填入无对应字段的要求。
    assert sheet.cell(2, len(headers)).value == '甲院启用移动报销'
    assert sheet.cell(2, len(headers) - 1).value == '材料.txt'


def test_unchanged_attachment_rows_are_not_returned_as_candidates(tmp_path):
    save(tmp_path, [item(attachments=[attachment(tmp_path)])])
    exported = api().export_qa(tmp_path, 'expense', confirmed=True)
    parsed = api().parse_qa_workbook(tmp_path, exported['path'])
    assert parsed['answers'] == [] and parsed['attachment_changes'] == []


def test_attempt_to_overwrite_version_workbook_is_rejected_by_parser(tmp_path):
    saved = save(tmp_path, [item()])
    book = load_workbook(saved['path'])
    book.worksheets[0].cell(load_template('expense')['questions'][1]['row'], answer_column(load_template('expense')) + 1, '不应原地编辑不可变版本')
    book.save(saved['path'])
    with pytest.raises(ValueError):
        api().parse_qa_workbook(tmp_path, saved['path'])


def test_missing_version_file_is_rejected_as_integrity_error(tmp_path):
    saved = save(tmp_path, [item()])
    (Path(saved['path']).parent / 'manifest.json').unlink()
    with pytest.raises(ValueError):
        api().read_qa(tmp_path, 'expense')


@pytest.mark.parametrize('field,value', [('scope', 'secret=123456'), ('confirmed_by', 'email=person@example.com')])
def test_sensitive_confirmation_context_is_not_persisted(tmp_path, field, value):
    with pytest.raises(ValueError, match='敏感'):
        save(tmp_path, [item()], **{field: value})
    assert list(tmp_path.iterdir()) == []


def test_blank_stale_answer_does_not_block_other_question(tmp_path):
    initial = save(tmp_path, [item()])
    old = records(initial)[item()['question_id']]['revision']
    save(tmp_path, [item(answer='聊天更新', revision=old)], batch='更新')
    result = save(tmp_path, [item(answer=' ', revision=old), item(index=1)], batch='空旧项与新题')
    assert result['answered'] == 2


def test_batch_receipt_failure_cannot_publish_version(tmp_path, monkeypatch):
    first = save(tmp_path, [item()])
    original = api().atomic_write
    def fail_receipt(path, content):
        if path.name == 'memory.json' and '失败批次' in content:
            raise OSError('批次确认写入失败')
        return original(path, content)
    monkeypatch.setattr(api(), 'atomic_write', fail_receipt)
    with pytest.raises(OSError, match='批次确认写入失败'):
        save(tmp_path, [item(index=1)], batch='失败批次')
    monkeypatch.undo()
    assert api().read_qa(tmp_path, 'expense')['answer_version'] == first['answer_version']
    assert save(tmp_path, [item(index=1)], batch='失败批次')['answered'] == 2


def test_symlinked_attachment_cannot_bind_external_material(tmp_path):
    external = tmp_path.parent / (tmp_path.name + '-material.txt')
    external.write_text('外部原材料')
    (tmp_path / '链接.txt').symlink_to(external)
    bound = {'reference': '链接.txt', 'sha256': hashlib.sha256(external.read_bytes()).hexdigest(), 'requirements': ['脱敏要求']}
    with pytest.raises(ValueError):
        save(tmp_path, [item(attachments=[bound])])


def test_scope_normalization_reuses_stable_identifier(tmp_path):
    initial = save(tmp_path, [item()], scope='甲院 / 本期')
    result = save(tmp_path, [item(index=1)], scope='甲院/本期', scope_id=initial['scope_id'], batch='规范范围')
    assert result['scope_id'] == initial['scope_id'] and result['answered'] == 2


def test_filtered_read_keeps_full_progress_and_does_not_write(tmp_path):
    save(tmp_path, [item(), item(index=1)])
    result = api().read_qa(tmp_path, 'expense', question_ids=[item()['question_id']])
    assert len(result['questions']) == len(result['records']) == 1
    assert result['answered'] == 2 and result['total'] == 25
    with pytest.raises(ValueError):
        api().read_qa(tmp_path, 'expense', question_ids=['不存在的问题'])


def test_metadata_module_and_revision_cannot_be_swapped(tmp_path):
    exported = api().export_qa(tmp_path, 'expense', confirmed=True)
    for field, value in [('module', 'general-ledger'), ('template_revision', '0' * 64), ('export_id', '../越界')]:
        book = load_workbook(exported['path'])
        original = json.loads(book['_dy_qa']['A1'].value)
        book['_dy_qa']['A1'] = json.dumps({**original, field: value})
        forged = tmp_path / f'{field}.xlsx'
        book.save(forged)
        with pytest.raises(ValueError):
            api().parse_qa_workbook(tmp_path, forged)


def test_blank_export_fill_and_commit_to_first_scope(tmp_path):
    exported = api().export_qa(tmp_path, 'expense', confirmed=True)
    book = load_workbook(exported['path'])
    assert book.worksheets[0].cell(load_template('expense')['questions'][0]['row'], answer_column(load_template('expense'))).value is None
    book.worksheets[0].cell(load_template('expense')['questions'][0]['row'], answer_column(load_template('expense')) + 1).value = '首次答卷答案'
    book.save(exported['path'])
    parsed = api().parse_qa_workbook(tmp_path, exported['path'])
    assert parsed['scope_id'] is None and parsed['base_version'] is None
    assert parsed['answers'][0]['research_complete'] is False
    confirmed_answers = [{**answer, 'research_complete': True} for answer in parsed['answers']]
    assert save(tmp_path, confirmed_answers, scope='首次确认范围')['answered'] == 1


def test_current_pointer_failure_retains_previous_version_and_retry_works(tmp_path, monkeypatch):
    initial = save(tmp_path, [item()])
    original = api().atomic_write
    def fail_pointer(path, content):
        if path.name == 'current.json':
            raise OSError('指针发布失败')
        return original(path, content)
    monkeypatch.setattr(api(), 'atomic_write', fail_pointer)
    with pytest.raises(OSError, match='指针发布失败'):
        save(tmp_path, [item(index=1)], batch='指针失败批次')
    monkeypatch.undo()
    assert api().read_qa(tmp_path, 'expense')['answer_version'] == initial['answer_version']
    assert save(tmp_path, [item(index=1)], batch='指针失败批次')['answered'] == 2


def test_semantically_inconsistent_xlsx_is_rejected_even_after_rehash(tmp_path):
    saved = save(tmp_path, [item()])
    directory = Path(saved['path']).parent
    book = load_workbook(saved['path'])
    book.worksheets[0].cell(load_template('expense')['questions'][0]['row'], answer_column(load_template('expense'))).value = '与机器记忆不一致'
    book.save(saved['path'])
    manifest_path = directory / 'manifest.json'
    envelope = json.loads(manifest_path.read_text())
    envelope['data']['hashes']['answers.xlsx'] = hashlib.sha256(Path(saved['path']).read_bytes()).hexdigest()
    canonical = json.dumps(envelope['data'], ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode()
    envelope['sha256'] = hashlib.sha256(canonical).hexdigest()
    manifest_path.write_text(json.dumps(envelope, ensure_ascii=False))
    pointer_path = directory.parent.parent / 'current.json'
    pointer = json.loads(pointer_path.read_text())
    pointer['manifest_sha256'] = hashlib.sha256(manifest_path.read_bytes()).hexdigest()
    pointer_path.write_text(json.dumps(pointer))
    with pytest.raises(ValueError, match='记忆'):
        api().read_qa(tmp_path, 'expense')


def test_batch_identifier_cannot_persist_sensitive_values(tmp_path):
    with pytest.raises(ValueError, match='敏感'):
        save(tmp_path, [item()], batch='token=不可保存的令牌')
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize('source', ['export', 'version_copy'])
def test_parse_preserves_partial_confirmation_without_upgrading_repeated_text(tmp_path, source):
    saved = save(tmp_path, [item(answer='待确认：仅部分材料', research_complete=False)])
    if source == 'export':
        path = Path(api().export_qa(tmp_path, 'expense', confirmed=True)['path'])
    else:
        path = tmp_path / '同步版回填.xlsx'
        path.write_bytes(Path(saved['path']).read_bytes())
    book = load_workbook(path)
    book.worksheets[0].cell(load_template('expense')['questions'][0]['row'], answer_column(load_template('expense')) + 1).value = '待确认：仅部分材料'
    book.save(path)
    parsed = api().parse_qa_workbook(tmp_path, path)
    assert parsed['answers'][0]['research_complete'] is False
    unchanged = save(tmp_path, parsed['answers'], batch='局部原文回填')
    assert unchanged['answer_version'] == saved['answer_version']
    assert unchanged['answered'] == 0 and unchanged['confirmed_count'] == 1


def test_parse_new_answer_defaults_to_partial_until_explicit_confirmation(tmp_path):
    path = Path(api().export_qa(tmp_path, 'expense', confirmed=True)['path'])
    book = load_workbook(path)
    book.worksheets[0].cell(load_template('expense')['questions'][0]['row'], answer_column(load_template('expense')) + 1).value = '待确认：仅部分材料'
    book.save(path)
    parsed = api().parse_qa_workbook(tmp_path, path)
    assert parsed['answers'][0]['research_complete'] is False
    partial = save(tmp_path, parsed['answers'], batch='确认局部材料')
    assert partial['answered'] == 0 and partial['confirmed_count'] == 1


def test_attachment_without_template_sheet_is_exported_and_change_is_candidate(tmp_path):
    template = load_template('general-ledger')
    qid = 'GL-2-009'
    index = next(i for i, q in enumerate(template['questions']) if q['question_id'] == qid)
    assert template['questions'][index]['attachment_names'] == []
    bound = attachment(tmp_path)
    save(tmp_path, [item(module='general-ledger', index=index, attachments=[bound])], module='general-ledger')
    exported = api().export_qa(tmp_path, 'general-ledger', confirmed=True)
    book = load_workbook(exported['path'])
    assert '已确认材料' in book.sheetnames
    sheet = book['已确认材料']
    assert list(sheet.values) == [('配置项', '调研问题', '材料名称', '已确认要求'),
                                  (template['questions'][index]['config_item'], '/', '材料.txt', '甲院启用移动报销')]
    sheet['D2'] = '候选：总账样张需补充审核要求'
    book.save(exported['path'])
    parsed = api().parse_qa_workbook(tmp_path, exported['path'])
    assert parsed['attachment_changes'][0]['question_ids'] == [qid]
    assert parsed['attachment_changes'][0]['can_write_confirmed'] is False


def test_export_keeps_real_scope_when_another_scope_appears_before_lock(tmp_path, monkeypatch):
    from contextlib import contextmanager
    first = save(tmp_path, [item(answer='甲院已确认答案')])
    original_lock = api().record_lock
    inserted = False

    @contextmanager
    def interleave(path):
        nonlocal inserted
        if not inserted:
            inserted = True
            with monkeypatch.context() as context:
                context.setattr(api(), 'record_lock', original_lock)
                save(tmp_path, [item(answer='乙院已确认答案')], scope='乙院/本期', batch='并发第二范围')
        with original_lock(path):
            yield

    monkeypatch.setattr(api(), 'record_lock', interleave)
    exported = api().export_qa(tmp_path, 'expense', confirmed=True)
    files = exported.get('files', [exported])
    assert all(file['scope_id'] is not None for file in files)
    assert first['scope_id'] in {file['scope_id'] for file in files}
    answers = {load_workbook(file['path']).worksheets[0].cell(load_template('expense')['questions'][0]['row'], answer_column(load_template('expense'))).value for file in files}
    assert '甲院已确认答案' in answers and None not in answers


def test_successful_batch_replay_after_attachment_drift_returns_receipt_and_warning(tmp_path):
    bound = attachment(tmp_path)
    request = [item(attachments=[bound])]
    first = save(tmp_path, request, batch='带附件成功批次')
    attachment(tmp_path, '后续材料内容已经变化')
    replay = save(tmp_path, request, batch='带附件成功批次')
    assert replay['idempotent'] is True and replay['stored'] is False
    assert replay['answer_version'] == first['answer_version']
    assert replay['attachment_integrity'] is False
    assert replay['attachment_changes'][0]['reference'] == '材料.txt'
    with pytest.raises(ValueError, match='摘要'):
        save(tmp_path, [item(index=1, attachments=[bound])], batch='新批次旧附件')
    with pytest.raises(ValueError, match='批次'):
        save(tmp_path, [item(answer='改变原请求', attachments=[bound])], batch='带附件成功批次')


def test_successful_replay_without_material_file_is_still_idempotent(tmp_path):
    bound = attachment(tmp_path)
    request = [item(attachments=[bound])]
    saved = save(tmp_path, request, batch='材料已确认批次')
    (tmp_path / '材料.txt').unlink()
    replay = save(tmp_path, request, batch='材料已确认批次')
    assert replay['idempotent'] is True
    assert replay['answer_version'] == saved['answer_version']
    assert replay['attachment_integrity'] is False


def test_blank_export_race_returns_separate_scope_files(tmp_path, monkeypatch):
    from contextlib import contextmanager
    original_lock = api().record_lock
    inserted = False

    @contextmanager
    def interleave(path):
        nonlocal inserted
        if not inserted:
            inserted = True
            with monkeypatch.context() as context:
                context.setattr(api(), 'record_lock', original_lock)
                save(tmp_path, [item(answer='甲院答案')], batch='首次甲院')
                save(tmp_path, [item(answer='乙院答案')], scope='乙院', batch='首次乙院')
        with original_lock(path):
            yield

    monkeypatch.setattr(api(), 'record_lock', interleave)
    exported = api().export_qa(tmp_path, 'expense', confirmed=True)
    assert len(exported['files']) == 2
    assert all(file['scope_id'] is not None for file in exported['files'])
    assert {load_workbook(file['path']).worksheets[0].cell(load_template('expense')['questions'][0]['row'], answer_column(load_template('expense'))).value for file in exported['files']} == {'甲院答案', '乙院答案'}


def test_parse_preserves_explicitly_complete_existing_answer(tmp_path):
    saved = save(tmp_path, [item(research_complete=True)])
    path = Path(api().export_qa(tmp_path, 'expense', confirmed=True)['path'])
    book = load_workbook(path)
    book.worksheets[0].cell(load_template('expense')['questions'][0]['row'], answer_column(load_template('expense')) + 1).value = '本期交付移动报销'
    book.save(path)
    parsed = api().parse_qa_workbook(tmp_path, path)
    assert parsed['answers'][0]['research_complete'] is True
    assert save(tmp_path, parsed['answers'], batch='完整原文回填')['answer_version'] == saved['answer_version']
