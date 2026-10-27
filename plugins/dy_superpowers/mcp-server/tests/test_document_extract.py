"""完整文档提取的内容保留与缺口契约。"""
import io
import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parents[1] / 'src'))


def extract(path):
    from dy_superpowers_mcp.document_extract import extract_document
    return extract_document(path)


def contents(result):
    return '\n'.join(unit['content'] for unit in result['units'])


def test_text_preserves_short_lines_code_tables_and_long_tail(tmp_path):
    source = '| 列 | 值 |\n| -- | -- |\n| A | B |\n```python\nx=1\n```\n短\n' + '正文' * 20000 + '\n末尾'
    path = tmp_path / '全部.md'
    path.write_text(source, encoding='utf-8')
    result = extract(path)
    assert contents(result) == source
    assert result['inventory']['complete'] is True
    assert result['issues'] == []


def test_json_retains_full_original_content(tmp_path):
    path = tmp_path / '全部.json'
    source = '{"a": [1, 2], "短": "值", "table": [["x"], ["y"]]}'
    path.write_text(source, encoding='utf-8')
    assert contents(extract(path)) == source


@pytest.mark.parametrize(('encoding', 'bom'), [
    ('utf-32-le', b'\xff\xfe\x00\x00'),
    ('utf-32-be', b'\x00\x00\xfe\xff'),
    ('utf-16-le', b'\xff\xfe'),
    ('utf-16-be', b'\xfe\xff'),
])
def test_text_decodes_unicode_bom_without_corrupting_content(tmp_path, encoding, bom):
    expected = '中文与非基本平面字符：𠮷 😀\n第二行'
    path = tmp_path / 'utf32.txt'
    path.write_bytes(bom + expected.encode(encoding))

    result = extract(path)

    assert contents(result) == expected
    assert result['inventory']['complete'] is True


def test_xlsx_keeps_all_sheets_hidden_cells_formula_comments_links(tmp_path):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    book = Workbook()
    book.active.title = '无表头'
    book.active['A1'] = '短'
    book.active['Z50'] = '尾值'
    book.active['B2'] = '=1+2'
    book.active['A1'].comment = Comment('补充批注', '作者')
    book.active['D4'].hyperlink = 'https://example.invalid/参考'
    book.active.merge_cells('A7:C7')
    hidden = book.create_sheet('隐藏页')
    hidden.sheet_state = 'hidden'
    hidden['A1'] = '隐藏数据'
    hidden.row_dimensions[1].hidden = True
    hidden.column_dimensions['A'].hidden = True
    path = tmp_path / '全部.xlsx'
    book.save(path)
    result = extract(path)
    text = contents(result)
    for expected in ['短', '尾值', '=1+2', '补充批注', 'https://example.invalid/参考', '隐藏数据', 'A7:C7']:
        assert expected in text
    assert result['inventory']['sheets'] == 2
    formula = next(unit for unit in result['units'] if unit.get('formula') == '=1+2')
    assert formula['cached_value'] is None
    assert any(issue['code'] == 'formula_cache_missing' for issue in result['issues'])
    hidden_unit = next(unit for unit in result['units'] if unit['content'] == '隐藏数据')
    assert hidden_unit['sheet_state'] == 'hidden'
    assert hidden_unit['row_hidden'] and hidden_unit['column_hidden']
    assert not result['inventory']['complete']


def test_docx_keeps_every_short_paragraph_nested_table_and_extra_parts(tmp_path):
    from docx import Document
    document = Document()
    document.add_paragraph('甲')
    document.add_paragraph('乙')
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = '表格值'
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = '嵌套值'
    document.sections[0].header.paragraphs[0].text = '页眉短文'
    path = tmp_path / '全部.docx'
    document.save(path)
    with zipfile.ZipFile(path, 'a') as package:
        package.writestr('word/footnotes.xml', '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>脚注值</w:t></w:r></w:p></w:footnote></w:footnotes>')
        package.writestr('word/comments.xml', '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:comment w:id="2"><w:p><w:del><w:r><w:delText>删除修订</w:delText></w:r></w:del><w:r><w:t>批注值</w:t></w:r></w:p></w:comment></w:comments>')
    result = extract(path)
    text = contents(result)
    for expected in ['甲', '乙', '表格值', '嵌套值', '页眉短文', '脚注值', '删除修订', '批注值']:
        assert expected in text
    assert result['inventory']['tables'] >= 2
    assert any(unit.get('revision') == 'deleted' for unit in result['units'])


def test_pdf_mixed_image_page_is_not_complete_when_ocr_missing(tmp_path, monkeypatch):
    from PIL import Image
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.utils import ImageReader
    from dy_superpowers_mcp import document_extract
    monkeypatch.setattr(document_extract.shutil, 'which', lambda name: None)
    image = Image.new('RGB', (80, 30), 'white')
    path = tmp_path / '混合.pdf'
    canvas = Canvas(str(path))
    canvas.drawString(20, 800, 'TEXT CONTENT')
    canvas.drawImage(ImageReader(image), 20, 700, width=80, height=30)
    canvas.showPage()
    canvas.drawString(20, 800, 'SECOND PAGE')
    canvas.save()
    result = extract(path)
    assert 'TEXT CONTENT' in contents(result)
    assert 'SECOND PAGE' in contents(result)
    assert result['inventory']['pages'] == 2
    assert result['assets']
    assert any('/image:' in asset['locator'] for asset in result['assets'])
    assert any(issue['code'] == 'ocr_unavailable' and 'page:1' in issue['locator'] for issue in result['issues'])
    assert not result['inventory']['complete']


def test_pptx_keeps_hidden_slides_notes_tables_group_text(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide._element.set('show', '0')
    slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = '短文'
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = '组内值'
    slide.shapes.add_table(1, 1, 0, 0, Inches(2), Inches(1)).table.cell(0, 0).text = '表格值'
    slide.notes_slide.notes_text_frame.text = '备注值'
    path = tmp_path / '全部.pptx'
    presentation.save(path)
    result = extract(path)
    for expected in ['短文', '组内值', '表格值', '备注值']:
        assert expected in contents(result)
    assert result['inventory']['slides'] == 1
    assert any(unit.get('hidden') is True for unit in result['units'])


def test_xlsx_chart_and_picture_are_preserved_with_review_gaps(tmp_path):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.drawing.image import Image as SpreadsheetImage
    from PIL import Image
    book = Workbook()
    sheet = book.active
    sheet['A1'] = '类别'
    sheet['B1'] = '数量'
    sheet['A2'] = '甲'
    sheet['B2'] = 5
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=2))
    sheet.add_chart(chart, 'D1')
    image_path = tmp_path / 'image.png'
    Image.new('RGB', (10, 10), 'white').save(image_path)
    sheet.add_image(SpreadsheetImage(str(image_path)), 'D10')
    path = tmp_path / '绘图.xlsx'
    book.save(path)
    result = extract(path)
    assert result['inventory']['charts'] == 1
    assert any(asset['name'].endswith('.png') for asset in result['assets'])
    assert any(issue['code'] == 'chart_visual_review_required' for issue in result['issues'])
    assert any(issue['code'] == 'visual_review_required' for issue in result['issues'])


def test_image_ocr_when_chinese_language_available_still_requires_visual_review(tmp_path, monkeypatch):
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen.canvas import Canvas
    from PIL import Image
    from dy_superpowers_mcp import document_extract
    image = Image.new('RGB', (80, 30), 'white')
    path = tmp_path / '需要人工核对.pdf'
    canvas = Canvas(str(path))
    canvas.drawImage(ImageReader(image), 20, 700, width=80, height=30)
    canvas.save()
    monkeypatch.setattr(document_extract.shutil, 'which', lambda name: '/usr/bin/tesseract')
    calls = []

    class Completed:
        def __init__(self, stdout):
            self.stdout = stdout

    def run(args, **kwargs):
        calls.append(args)
        return Completed(b'List of available languages (2):\nchi_sim\neng\n' if len(calls) == 1 else '识别文字'.encode())

    monkeypatch.setattr(document_extract.subprocess, 'run', run)
    result = extract(path)
    assert any(unit['kind'] == 'ocr_text' and unit['content'] == '识别文字' for unit in result['units'])
    assert any(issue['code'] == 'visual_review_required' for issue in result['issues'])
    assert not result['inventory']['complete']


def test_pptx_chart_series_are_extracted_without_losing_visual_gap(tmp_path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ['甲', '乙']
    data.add_series('数量', [2, 4])
    slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, 0, 0, Inches(4), Inches(3), data)
    path = tmp_path / '图表.pptx'
    presentation.save(path)
    result = extract(path)
    assert result['inventory']['charts'] == 1
    assert '甲' in contents(result) and '乙' in contents(result)
    assert '数量' in contents(result)
    assert any(issue['code'] == 'chart_visual_review_required' for issue in result['issues'])


def test_xlsx_markdown_conversion_is_a_separate_derived_asset(tmp_path, monkeypatch):
    import types
    from openpyxl import Workbook
    book = Workbook()
    book.active['B9'] = '=1+2'
    path = tmp_path / 'sheet.xlsx'
    book.save(path)

    class Converted:
        text_content = '# Readable view\n\n| value |\n|---|'

    class MarkItDown:
        def __init__(self, **kwargs):
            assert kwargs == {'enable_plugins': False}

        def convert(self, source):
            assert Path(source) == path
            return Converted()

    monkeypatch.setitem(sys.modules, 'markitdown', types.SimpleNamespace(MarkItDown=MarkItDown))
    result = extract(path)
    derived = next(asset for asset in result['assets'] if asset['name'] == 'sheet.converted.md')
    assert derived['data'].decode('utf-8') == Converted.text_content
    assert any(unit['kind'] == 'converted_markdown' and 'Readable view' in unit['content'] for unit in result['units'])
    assert any(unit.get('formula') == '=1+2' for unit in result['units'])


def test_markdown_conversion_failure_keeps_structured_parse_and_blocks_completeness(tmp_path, monkeypatch):
    import types
    from openpyxl import Workbook
    book = Workbook()
    book.active['A1'] = '保留值'
    path = tmp_path / 'sheet.xlsx'
    book.save(path)

    class MarkItDown:
        def __init__(self, **kwargs):
            assert kwargs == {'enable_plugins': False}

        def convert(self, source):
            raise RuntimeError('converter failed')

    monkeypatch.setitem(sys.modules, 'markitdown', types.SimpleNamespace(MarkItDown=MarkItDown))
    result = extract(path)
    assert '保留值' in contents(result)
    assert any(issue['code'] == 'markdown_conversion_failed' for issue in result['issues'])
    assert not result['inventory']['complete']


@pytest.mark.parametrize('suffix', ['xlsx', 'docx'])
def test_markitdown_real_conversion_creates_readable_view_separate_from_source_parse(tmp_path, suffix):
    if suffix == 'xlsx':
        from openpyxl import Workbook
        document = Workbook()
        document.active['A1'] = '类目'
        document.active['B1'] = '值'
        document.active['A2'] = '收入'
        document.active['B2'] = 12
        document.active['B3'] = '=B2*2'
    else:
        from docx import Document
        document = Document()
        document.add_paragraph('短段落')
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = '字段'
        table.cell(0, 1).text = '说明'
        row = table.add_row()
        row.cells[0].text = 'A'
        row.cells[1].text = '值'
    path = tmp_path / f'视图.{suffix}'
    document.save(path)
    result = extract(path)
    converted = next(unit['content'] for unit in result['units'] if unit['kind'] == 'converted_markdown')
    asset = next(asset for asset in result['assets'] if asset.get('role') == 'derived_markdown')
    assert asset['data'].decode('utf-8') == converted
    assert converted.strip()
    assert asset['formatter'] == 'markitdown'
    if suffix == 'xlsx':
        assert any(unit.get('coordinate') == 'B3' and unit.get('formula') == '=B2*2' for unit in result['units'])
        assert any(issue['code'] == 'formula_cache_missing' for issue in result['issues'])
        assert not result['inventory']['complete']
    else:
        assert '短段落' in contents(result)
        assert any(unit['kind'] == 'table' for unit in result['units'])


@pytest.mark.parametrize('filename,data,code', [
    ('旧版.doc', b'legacy binary', 'unsupported_format'),
    ('损坏.xlsx', b'bad zip', 'parse_error'),
    ('损坏.pdf', b'bad pdf', 'parse_error'),
    ('未知.bin', b'abc', 'unsupported_format'),
])
def test_unreadable_formats_have_explicit_blocking_issue(tmp_path, filename, data, code):
    path = tmp_path / filename
    path.write_bytes(data)
    result = extract(path)
    assert any(issue['code'] == code for issue in result['issues'])
    assert result['inventory']['complete'] is False
