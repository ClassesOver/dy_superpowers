"""真实 CLI 来源矩阵：候选、全文与附件缺口不得互相掩盖。"""
import json
import os
from pathlib import Path
import subprocess
import sys

import pytest

SCRIPT = Path(__file__).parents[2] / 'skills/module-research/scripts/research_sources.py'


def cli(workspace, source, *extra):
    env = {k: v for k, v in os.environ.items()
           if not k.startswith(('OTEL_', 'DY_WIKI_', 'DY_SUPERPOWERS_OTLP_'))}
    env['CODEX_HOME'] = str(workspace / 'codex')
    env['PATH'] = str(Path(sys.executable).parent)
    return subprocess.run(
        [sys.executable, str(SCRIPT), str(source), '--module', 'general-ledger',
         '--workspace', str(workspace), *extra],
        capture_output=True, text=True, timeout=45, env=env,
    )


def parsed(workspace, source, *extra):
    completed = cli(workspace, source, *extra)
    assert completed.returncode == 0, completed.stderr
    result = json.loads(completed.stdout)
    run = workspace / result['path']
    return result, {name: json.loads((run / (name + '.json')).read_text())
                    for name in ('content', 'candidates', 'coverage', 'issues', 'manifest')}


def qa(path, attachments=(), answer='本期基础应用必需'):
    path.write_text(json.dumps({'items': [{'question': '本期应用范围？',
                                          'answer': answer, 'scope': '甲医院/本期',
                                          'attachments': list(attachments)}]},
                               ensure_ascii=False), encoding='utf-8')
    return path


def make_document(path):
    if path.suffix == '.xlsx':
        from openpyxl import Workbook
        book = Workbook()
        book.active.append(['检查编号', '配置项', '调研结论'])
        book.active.append(['GL-1-001', '应用安装', '01｜总账基础应用本期必需'])
        book.save(path)
    elif path.suffix == '.docx':
        from docx import Document
        document = Document()
        document.add_paragraph('总账基础应用调研结论：本期必需')
        document.add_table(rows=1, cols=1).cell(0, 0).text = '表格尾部内容'
        document.save(path)
    elif path.suffix == '.pptx':
        from pptx import Presentation
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide._element.set('show', '0')
        slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = '幻灯片正文'
        slide.notes_slide.notes_text_frame.text = '备注尾部内容'
        presentation.save(path)
    elif path.suffix == '.pdf':
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                 NameObject('/Subtype'): NameObject('/Type1'),
                                 NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):
            DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 20 250 Td (RESULT: REQUIRED) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
        writer.write(path)
    elif path.suffix == '.png':
        from PIL import Image
        Image.new('RGB', (30, 20), 'white').save(path)
    else:
        path.write_text('总账基础应用本期必需；尾部𠮷😀', encoding='utf-8')
    return path


@pytest.mark.parametrize('suffix, marker', [
    ('.xlsx', '01｜总账基础应用本期必需'), ('.docx', '表格尾部内容'),
    ('.pdf', 'RESULT: REQUIRED'), ('.txt', '尾部𠮷😀'),
])
def test_cli_preserves_local_source_and_reuses_complete_run(tmp_path, suffix, marker):
    source = make_document(tmp_path / ('本地' + suffix))
    first, data = parsed(tmp_path, source)
    assert first['status'] == 'complete'
    assert marker in json.dumps(data['content'], ensure_ascii=False)
    assert not (tmp_path / '.dy_superpowers').exists()
    second, _ = parsed(tmp_path, source)
    assert second['path'] == first['path']
    forced, _ = parsed(tmp_path, source, '--force')
    assert forced['path'] != first['path']
    assert (tmp_path / first['path'] / 'manifest.json').exists()


@pytest.mark.parametrize('suffix', ['.xlsx', '.docx', '.pdf', '.txt'])
def test_local_attachment_archives_without_expanding_answer(tmp_path, suffix):
    child = make_document(tmp_path / ('明细' + suffix))
    result, data = parsed(tmp_path, qa(tmp_path / 'QA.json', [child.name]))
    assert result['status'] == 'complete'
    assert len(result['attachments']) == 1
    attachment = result['attachments'][0]
    assert attachment['status'] == 'complete'
    assert (tmp_path / attachment['path'] / 'content.json').exists()
    record = data['candidates']['records'][0]
    assert record['conclusion'] == '本期基础应用必需'
    assert record['can_write_confirmed'] is False
    assert record['detail_refs'][0]['archive_path'] == attachment['path']
    assert 'sample' not in record['detail_refs'][0]


@pytest.mark.parametrize('encoding', ['utf-8-sig', 'utf-16-le', 'utf-16-be', 'utf-32-le', 'utf-32-be'])
def test_cli_unicode_candidate_and_full_text_agree(tmp_path, encoding):
    text = '问：本期应用？\n答：甲医院𠮷😀本期基础应用必需'
    bom = {'utf-16-le': b'\xff\xfe', 'utf-16-be': b'\xfe\xff',
           'utf-32-le': b'\xff\xfe\x00\x00', 'utf-32-be': b'\x00\x00\xfe\xff'}.get(encoding, b'')
    source = tmp_path / 'Unicode.txt'
    source.write_bytes(bom + text.encode(encoding))
    result, data = parsed(tmp_path, source)
    assert result['status'] == 'complete'
    assert [r['conclusion'] for r in data['candidates']['records']] == ['甲医院𠮷😀本期基础应用必需']
    assert any(u['content'] == text for u in data['content']['units'])


@pytest.mark.parametrize('suffix', ['.pptx', '.png'])
def test_cli_missing_candidate_reader_still_reports_structured_status(tmp_path, suffix):
    source = make_document(tmp_path / ('来源' + suffix))
    result, data = parsed(tmp_path, source)
    assert result['status'] in {'partial', 'failed'}
    assert any(issue.get('blocking', True) for issue in data['issues'])


def test_cli_xlsx_candidate_and_document_gaps_keep_hidden_merged_content(tmp_path):
    from openpyxl import Workbook
    book = Workbook(); sheet = book.active; sheet.title = '主表'
    sheet.append(['检查编号', '配置项', '调研结论'])
    sheet.append(['GL-1-001', '应用安装', '01｜本期基础应用必需'])
    sheet.append([None, None, '02｜扩展应用待确认'])
    sheet.merge_cells('A2:A3'); sheet.merge_cells('B2:B3')
    sheet['C4'] = '=C2'; sheet.row_dimensions[3].hidden = True
    hidden = book.create_sheet('隐藏'); hidden.sheet_state = 'hidden'; hidden['A1'] = '隐藏尾部'
    source = tmp_path / '复杂.xlsx'; book.save(source)
    result, data = parsed(tmp_path, source)
    assert result['status'] == 'partial'
    assert {'formula_without_cache', 'formula_cache_missing', 'hidden_sheet'} <= {x['code'] for x in data['issues']}
    rows = data['candidates']['records']
    assert [(r['check_order'], r['research_id']) for r in rows] == [('GL-1-001', '01'), ('GL-1-001', '02')]
    assert rows[1]['hidden'] is True
    assert '隐藏尾部' in json.dumps(data['content'], ensure_ascii=False)
    assert data['coverage']['unresolved'] == result['issues']
    assert any(u['kind'] == 'merged_range' for u in data['content']['units'])


@pytest.mark.parametrize('missing_column', [0, 1, 2])
def test_cli_blank_required_xlsx_field_blocks_complete(tmp_path, missing_column):
    from openpyxl import Workbook
    book = Workbook(); book.active.append(['检查编号', '配置项', '调研结论'])
    row = ['GL-1-001', '应用安装', '01｜本期基础应用必需']; row[missing_column] = None
    book.active.append(row); source = tmp_path / '缺关键字段.xlsx'; book.save(source)
    result, data = parsed(tmp_path, source)
    assert result['status'] == 'partial'
    assert any(x.get('source_component') == 'candidate_parser' for x in data['issues'])


@pytest.mark.parametrize('answer', ['', '   ', None])
def test_cli_blank_qa_answer_stays_blocked(tmp_path, answer):
    result, data = parsed(tmp_path, qa(tmp_path / '缺答.json', answer=answer))
    assert result['status'] == 'partial' and result['source_subitems'] == 0
    assert any(x['code'] == 'qa_item_missing_question_or_answer' for x in data['issues'])



@pytest.mark.parametrize('kind, expected_issue', [
    ('missing', 'attachment_missing'), ('corrupt', 'attachment_parse_failed'),
    ('unsupported', 'attachment_unsupported_type'), ('credentials', 'attachment_rejected'),
])
def test_cli_attachment_failures_block_main(tmp_path, kind, expected_issue):
    names = {'missing': '不存在.txt', 'corrupt': '损坏.docx',
             'unsupported': '旧版.xls', 'credentials': 'https://user:synthetic-secret@example.invalid/a'}
    if kind in {'corrupt', 'unsupported'}:
        (tmp_path / names[kind]).write_bytes(b'invalid container')
    result, data = parsed(tmp_path, qa(tmp_path / '附件状态.json', [names[kind]]))
    assert result['status'] == 'partial'
    assert any(x['code'] == expected_issue for x in data['issues'])
    assert 'synthetic-secret' not in json.dumps(result)


def test_cli_remote_reference_never_opens_socket_or_downloads(tmp_path):
    import socket
    with socket.socket() as listener:
        listener.bind(('127.0.0.1', 0)); listener.listen(); listener.settimeout(0.1)
        url = f'http://127.0.0.1:{listener.getsockname()[1]}/private.txt'
        result, data = parsed(tmp_path, qa(tmp_path / '远程.json', [url]))
        with pytest.raises(TimeoutError):
            listener.accept()
    assert result['attachments'] == [{'name': 'private.txt', 'status': 'not_local', 'path': None}]
    assert data['candidates']['records'][0]['detail_refs'][0]['fetched'] is False
    assert not (tmp_path / 'private.txt').exists()


def test_cli_duplicate_attachment_reuses_archive_without_losing_occurrences(tmp_path):
    child = make_document(tmp_path / '共同.txt')
    first, data = parsed(tmp_path, qa(tmp_path / '重复.json', [child.name, child.name]))
    assert len(first['attachments']) == 2
    assert first['attachments'][0]['path'] == first['attachments'][1]['path']
    assert len(data['candidates']['records'][0]['detail_refs']) == 2
    second, _ = parsed(tmp_path, tmp_path / '重复.json')
    assert second['path'] == first['path']


def test_cli_complete_main_with_partial_child_cannot_report_complete(tmp_path):
    from docx import Document
    from PIL import Image
    image = tmp_path / 'image.png'; Image.new('RGB', (20, 20), 'white').save(image)
    document = Document(); document.add_paragraph('附件正文：预算已启用，只是观察')
    document.add_picture(str(image)); child = tmp_path / '视觉附件.docx'; document.save(child)
    result, data = parsed(tmp_path, qa(tmp_path / '主档.json', [child.name]))
    assert result['attachments'][0]['status'] == 'partial'
    assert result['status'] == 'partial', '附件含视觉/OCR阻塞，主档不得汇总为complete'
    assert result['issues'] > 0
    assert data['candidates']['records'][0]['conclusion'] == '本期基础应用必需'


def test_cli_candidate_document_attachment_issues_all_enter_summary(tmp_path):
    from openpyxl import Workbook
    book = Workbook(); book.active['A1'] = '=1+2'; child = tmp_path / '缺缓存.xlsx'; book.save(child)
    source = qa(tmp_path / '联动.json', [child.name, '缺失.txt'])
    payload = json.loads(source.read_text())
    payload['items'].append({'question': '另一个问题', 'answer': ''})
    source.write_text(json.dumps(payload, ensure_ascii=False))
    result, data = parsed(tmp_path, source)
    assert result['status'] == 'partial'
    assert {'qa_item_missing_question_or_answer', 'attachment_missing'} <= {x['code'] for x in data['issues']}
    child_info = next(x for x in result['attachments'] if x['name'] == child.name)
    child_issues = json.loads((tmp_path / child_info['path'] / 'issues.json').read_text())
    assert any(x['code'] == 'formula_cache_missing' for x in child_issues)
    assert result['issues'] >= 3, '候选、文档和附件问题都应进入汇总计数'


def test_cli_sensitive_values_do_not_leak_to_stdout_or_text_archive(tmp_path):
    source = qa(tmp_path / '脱敏.json', answer='password: synthetic-secret 手机号：13800138000 基础应用必需')
    completed = cli(tmp_path, source); assert completed.returncode == 0
    result = json.loads(completed.stdout); run = tmp_path / result['path']
    output = completed.stdout + completed.stderr + ''.join(p.read_text() for p in run.glob('*.json')) + (run / 'content.md').read_text()
    assert 'synthetic-secret' not in output and '13800138000' not in output
    assert '[REDACTED]' in output and '基础应用必需' in output


@pytest.mark.parametrize('suffix', ['.xlsx', '.docx', '.pdf', '.xls'])
def test_cli_corrupt_primary_returns_structured_failure(tmp_path, suffix):
    source = tmp_path / ('不可读' + suffix); source.write_bytes(b'not a valid document')
    result, data = parsed(tmp_path, source)
    assert result['status'] in {'partial', 'failed'}
    assert data['issues']
    assert not data['coverage']['inventory']['complete']



@pytest.mark.parametrize('suffix', ['.pptx', '.png'])
def test_cli_local_visual_attachment_gets_an_archive_not_not_local(tmp_path, suffix):
    child = make_document(tmp_path / ('本地视觉' + suffix))
    result, _ = parsed(tmp_path, qa(tmp_path / '视觉附件.json', [child.name]))
    attachment = result['attachments'][0]
    assert attachment['path'] is not None, '实际本地附件不能误报为not_local且不归档'
    assert attachment['status'] in {'complete', 'partial', 'failed'}
    assert result['status'] != 'complete' or attachment['status'] == 'complete'


def test_cli_corrupt_child_has_a_failed_or_partial_archive(tmp_path):
    child = tmp_path / '损坏子档.docx'; child.write_bytes(b'broken zip')
    result, data = parsed(tmp_path, qa(tmp_path / '子档失败.json', [child.name]))
    assert result['status'] == 'partial'
    attachment = result['attachments'][0]
    assert attachment['status'] in {'partial', 'failed'}
    assert attachment['path'] is not None
    child_coverage = json.loads((tmp_path / attachment['path'] / 'coverage.json').read_text())
    assert child_coverage['unresolved'] > 0
    assert data['candidates']['records'][0]['conclusion'] == '本期基础应用必需'


def test_cli_redaction_keeps_safe_business_text_after_secret(tmp_path):
    source = qa(tmp_path / '保留业务结论.json', answer='password: synthetic-secret；基础应用仍须上线')
    result, data = parsed(tmp_path, source)
    row = data['candidates']['records'][0]
    assert 'synthetic-secret' not in row['conclusion']
    assert '基础应用仍须上线' in row['conclusion'], '脱敏不得把同一回答后面的业务含义一并删除'
    assert result['status'] == 'complete'


def test_cli_sensitive_docx_text_is_redacted_in_derived_and_xml_assets(tmp_path):
    from docx import Document
    document = Document()
    document.add_paragraph('password: synthetic-asset-secret')
    document.add_paragraph('业务尾部仍保留')
    source = tmp_path / '敏感来源.docx'; document.save(source)
    result, _ = parsed(tmp_path, source)
    run = tmp_path / result['path']
    leaked = [p.relative_to(run).as_posix() for p in run.rglob('*')
              if p.is_file() and b'synthetic-asset-secret' in p.read_bytes()]
    assert leaked == [], f'敏感文字泄漏到归档资产：{leaked}'
    assert '业务尾部仍保留' in (run / 'content.md').read_text()
