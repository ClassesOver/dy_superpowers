"""调研闭环缺口回归；仅使用临时工作区，不修改项目业务记忆。"""
import importlib.util
import json
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).parents[1] / 'src'))
from dy_superpowers_mcp.qa_store import read_qa, export_qa, parse_qa_workbook
from dy_superpowers_mcp.qa_templates import load_template
from dy_superpowers_mcp.qa_workbook import answer_column
from template_qa_helpers import save_answers
import server
from dy_superpowers_mcp import research as research_lib


def test_mcp_forwards_only_required_business_environment_keys():
    config = json.loads((Path(__file__).parents[2] / '.mcp.json').read_text())
    forwarded = set(config['mcpServers']['dy-superpowers-odoo']['env_vars'])
    assert set(server.ENV_FILE_KEYS) <= forwarded
    assert {'DY_WORKSPACE_ROOT', 'FEISHU_CODEX_BOT_SENDER_USER_ID'} <= forwarded
    assert 'OPENAI_API_KEY' not in forwarded
    assert 'FEISHU_BOT_APP_SECRET' not in forwarded

script = Path(__file__).parents[2] / 'skills/module-research/scripts/research_sources.py'
spec = importlib.util.spec_from_file_location('repair_research_sources', script)
sources = importlib.util.module_from_spec(spec)
spec.loader.exec_module(sources)


def test_export_filled_return_preserves_question_revision_and_candidate_boundary(tmp_path):
    state = save_answers(tmp_path, 'general-ledger')
    result = export_qa(tmp_path, 'general-ledger', confirmed=True)
    book = load_workbook(result['path'])
    q = load_template('general-ledger')['questions'][1]
    book.active.cell(q['row'], answer_column(load_template('general-ledger')) + 1, '本期新回答')
    book.save(result['path']); book.close()
    parsed = parse_qa_workbook(tmp_path, result['path'])
    assert parsed['can_write_confirmed'] is False
    assert parsed['answers'][0]['question_id'] == q['question_id']
    assert parsed['answers'][0]['expected_revision'] is None
    assert read_qa(tmp_path, 'general-ledger')['answer_version'] == state['answer_version']


def test_detail_pending_count_is_separate_from_candidate_count(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.append(['检查顺序', '配置项', '调研结果'])
    sheet.append(['GL-1-001', '应用', '01｜已配置'])
    detail = book.create_sheet('期间明细')
    detail.append(['上线期间', '待确认'])
    path = tmp_path / 'result.xlsx'
    book.save(path)
    parsed = sources.parse_source(path, 'general-ledger')
    assert parsed['coverage']['contains_pending_questions'] == 0
    assert parsed['coverage']['source_pending_cells'] == 1
    assert parsed['pending_locations'] == [{'sheet': '期间明细', 'cell': 'B1', 'hidden': False}]


def test_partial_answers_are_visible_without_claiming_completion(tmp_path):
    state = save_answers(tmp_path, 'general-ledger', complete=False)
    assert state['confirmed_count'] == state['partial_count'] == 1
    assert state['answered'] == 0 and state['total'] == len(state['questions'])


def test_new_progress_does_not_read_legacy_pages(tmp_path, monkeypatch):
    monkeypatch.setenv('CODEX_HOME', str(tmp_path / 'codex'))
    total=len(load_template('general-ledger')['questions'])
    save_answers(tmp_path, 'general-ledger', count=total)
    monkeypatch.setattr(research_lib, 'read_module_records', lambda *a, **kw: pytest.fail('新版不得读取历史分页'))
    result, error = server.execute_tool('dy_research_progress',
        dict(trace_id='new-routing', module='general-ledger', intent='research'), tmp_path)
    assert error is None and result['answered'] == total and result['route'] == 'interactive'


def test_multiple_result_sheets_are_not_silently_discarded(tmp_path):
    book = Workbook()
    for sheet, key in [(book.active, 'GL-1-001'), (book.create_sheet('第二批'), 'GL-1-002')]:
        sheet.append(['检查顺序', '配置项', '调研结果'])
        sheet.append([key, '应用', '01｜已配置'])
    path = tmp_path / 'batches.xlsx'
    book.save(path)
    result = sources.parse_source(path, 'general-ledger')
    assert [row['check_order'] for row in result['records']] == ['GL-1-001', 'GL-1-002']
