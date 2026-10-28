"""候选入口修复的边界回归，补足组合矩阵的去重与兼容检查。"""
import importlib.util
import json
from pathlib import Path
import subprocess
import sys

import pytest

SCRIPT = Path(__file__).parents[2] / 'skills/module-research/scripts/research_sources.py'
spec = importlib.util.spec_from_file_location('research_sources_cli_recovery', SCRIPT)
sources = importlib.util.module_from_spec(spec)
spec.loader.exec_module(sources)


@pytest.mark.parametrize('encoding', ['utf-16', 'utf-32'])
def test_qa_json_unicode_bom_decodes_strictly(tmp_path, encoding):
    path = tmp_path / '问答.json'
    path.write_text(json.dumps({'items': [{'question': '本期范围？', 'answer': '甲𠮷医院'}]}, ensure_ascii=False), encoding=encoding)
    assert sources.parse_source(path, 'general-ledger')['records'][0]['conclusion'] == '甲𠮷医院'


def test_invalid_unicode_is_issue_instead_of_replacement_candidates(tmp_path):
    path = tmp_path / '坏编码.txt'; path.write_bytes(b'\xffabc')
    parsed = sources.parse_source(path, 'general-ledger')
    assert parsed['records'] == []
    assert any(issue['code'] == 'candidate_parse_failed' for issue in parsed['issues'])


def test_blank_template_rows_do_not_become_missing_required_issues(tmp_path):
    from openpyxl import Workbook
    book = Workbook(); sheet = book.active
    sheet.append(['检查编号', '配置项', '调研结论'])
    sheet.append(['GL-1-001', '应用安装', '01｜本期必需'])
    sheet.cell(8, 3).number_format = '@'
    path = tmp_path / '空模板.xlsx'; book.save(path)
    parsed = sources.parse_source(path, 'general-ledger')
    assert len(parsed['records']) == 1
    assert not any(issue['code'] == 'research_field_missing' for issue in parsed['issues'])


def test_duplicate_partial_child_is_archived_once_and_issues_not_multiplied(tmp_path):
    from openpyxl import Workbook
    book = Workbook(); book.active['A1'] = '=1+2'
    child = tmp_path / '公式.xlsx'; book.save(child)
    source = tmp_path / 'QA.json'
    source.write_text(json.dumps({'items': [{'question': '范围？', 'answer': '甲医院', 'attachments': [child.name, child.name]}]}, ensure_ascii=False))
    completed = subprocess.run([sys.executable, str(SCRIPT), str(source), '--module', 'general-ledger', '--workspace', str(tmp_path)], capture_output=True, text=True, timeout=40)
    assert completed.returncode == 0, completed.stderr
    result = json.loads(completed.stdout)
    assert result['status'] == 'partial'
    assert result['attachments'][0]['path'] == result['attachments'][1]['path']
    issues = json.loads((tmp_path / result['path'] / 'issues.json').read_text())
    formula_issues = [issue for issue in issues if issue['code'] == 'formula_cache_missing']
    assert len(formula_issues) == 1




def test_formula_candidate_keeps_cached_text_after_required_field_validation(tmp_path):
    import zipfile
    import xml.etree.ElementTree as ET
    from openpyxl import Workbook
    book = Workbook(); sheet = book.active
    sheet.append(['检查编号', '配置项', '调研结论'])
    sheet.append(['GL-1-001', '应用安装', '="01｜缓存确认结论"'])
    path = tmp_path / '有缓存.xlsx'; book.save(path)
    with zipfile.ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    xml = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
    cell = next(c for c in xml.iter('{'+ns+'}c') if c.get('r') == 'C2')
    cell.set('t', 'str')
    cell.find('{'+ns+'}v').text = '01｜缓存确认结论'
    parts['xl/worksheets/sheet1.xml'] = ET.tostring(xml)
    with zipfile.ZipFile(path, 'w') as package:
        for name, value in parts.items():
            package.writestr(name, value)
    parsed = sources.parse_source(path, 'general-ledger')
    assert parsed['records'][0]['conclusion'] == '01｜缓存确认结论'
    assert parsed['records'][0]['research_id'] == '01'
    assert parsed['issues'] == []
