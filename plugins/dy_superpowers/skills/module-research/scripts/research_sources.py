"""调研来源统一解析：xlsx/docx/pdf/md/txt → 候选结论 JSON（含来源坐标与语义分类）。

读取纪律迁移自本环境成熟文件处理技能（spreadsheets/documents/pdf）：合并单元格按左上角
展开、公式结论核对引用与缓存、隐藏行列清点、无法解析的区域显式列入 issues 而非跳过。
仅做只读提取与候选分类；业务确认、矩阵映射与记忆写入由 module-research 流程决定。
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path
from urllib.parse import urlsplit

MCP_SRC = Path(__file__).parents[3] / "mcp-server" / "src"
sys.path.insert(0, str(MCP_SRC))

from dy_superpowers_mcp.source_archive import archive_source

SUBITEM = re.compile(r"^\s*(\d{1,3})\s*[｜|．.、]\s*(.+)$", re.DOTALL)
PENDING = re.compile(r"待确认|需确认|尚未确认|未签认|待补充|请确认")
REQUIREMENT = re.compile(r"需存在|需配置|需完成|需维护|需要安装|需要启用|必须安装|需按|必须")
OBSERVATION = re.compile(r"已配置|已建|存在\d|配置\d|系统有|当前|未配置|未发现|条记录|均已过期")
DECISION = re.compile(r"不适用|不交付|不需要|不迁移|不对历史|本期不启用|不续期|无外币业务|无外币核算")
QUESTION = re.compile(r"^\s*(?:问|Q|问题|提问)\s*[:：]\s*(.+)$", re.IGNORECASE)
ANSWER = re.compile(r"^\s*(?:答|A|回答|答案)\s*[:：]\s*(.*)$", re.IGNORECASE)
ATTACHMENT = re.compile(r"^\s*(?:附件|链接|材料|attachment)\s*[:：]\s*(.+)$", re.IGNORECASE)


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").strip()


def _unicode_text(path: Path) -> str:
    raw = path.read_bytes()
    encoding = "utf-8-sig"
    if raw.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        encoding = "utf-32"
    elif raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        encoding = "utf-16"
    return raw.decode(encoding)


def _semantics(text: str) -> dict:
    kinds = [name for name, flag in
             (("requirement", REQUIREMENT.search(text)), ("observation", OBSERVATION.search(text)),
              ("decision", DECISION.search(text))) if flag]
    return {"statement_kinds": kinds or ["source_statement"],
            "contains_pending_questions": bool(PENDING.search(text)),
            "can_write_confirmed": False}


def _row(order, item, conclusion, locator, applicability="", evidence="", detail_refs=None, hidden=False):
    numbered = SUBITEM.match(conclusion)
    return {"check_order": _text(order), "config_item": _text(item),
            "research_id": numbered.group(1) if numbered else "",
            "conclusion": conclusion, "locator": locator,
            "applicability": _text(applicability), "suggested_evidence": _text(evidence),
            "detail_refs": detail_refs or [], "hidden": hidden, **_semantics(conclusion)}


def parse_xlsx(path: Path) -> dict:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    worksheets, issues, tables = [], [], []  # tables: (表头标签, 行集合)
    unanswered, pending_locations = [], []
    for ws in workbook:
        for cells in ws:
            for cell in cells:
                if cell.data_type != 'f' and PENDING.search(_text(cell.value)):
                    pending_locations.append({'sheet': ws.title, 'cell': cell.coordinate,
                                              'hidden': ws.sheet_state != 'visible' or bool(ws.row_dimensions[cell.row].hidden)})
        merges = [(a.min_row, a.max_row, a.min_col, a.max_col) for a in ws.merged_cells.ranges]
        worksheets.append({"sheet": ws.title, "visibility": ws.sheet_state, "rows": ws.max_row,
                           "columns": ws.max_column, "merged_ranges": len(merges),
                           "formula_cells": sum(c.data_type == "f" for r in ws for c in r),
                           "hidden_rows": sum(bool(v.hidden) for v in ws.row_dimensions.values())})
        if ws.sheet_state != "visible":
            issues.append({"code": "hidden_sheet", "sheet": ws.title})
        header = None
        for row in range(1, min(ws.max_row, 20) + 1):
            names = {_text(c.value): c.column for c in ws[row] if c.value is not None}
            conclusion = next((k for k in names if k.startswith(("调研结果", "调研结论"))), None)
            if conclusion is None and '答案' in names and '调研子项' in names:
                conclusion = '答案'
            order = next((k for k in ("检查编号", "检查顺序") if k in names), None)
            if conclusion and order and "配置项" in names:
                header = row, names, conclusion, order
                break
        if not header:
            continue
        header_row, names, conclusion_label, order_label = header
        table_rows = []

        def structural(row, col):
            if not col:
                return None
            value = ws.cell(row, col).value
            if value is not None:
                return value
            for low, high, left, right in merges:
                if low <= row <= high and left <= col <= right:
                    return ws.cell(low, left).value
            return None

        for row in range(header_row + 1, ws.max_row + 1):
            cell = ws.cell(row, names[conclusion_label])
            # 只有实际业务行才校验；格式化空行和普通附件模板不是调研行。
            if not any(_text(ws.cell(row, col).value) for col in names.values()):
                continue
            value = cell.value
            if conclusion_label == '答案' and not _text(value):
                unanswered.append({'check_order': _text(structural(row, names[order_label])),
                                   'research_id': _text(structural(row, names.get('调研子项'))),
                                   'locator': {'sheet': ws.title, 'cell': cell.coordinate}})
                continue
            if cell.data_type == "f":
                target = cached[ws.title][cell.coordinate].value
                if target is None:
                    issues.append({"code": "formula_without_cache", "sheet": ws.title, "cell": cell.coordinate})
                    continue
                value = target
            missing = False
            for label in (order_label, "配置项", conclusion_label):
                field_value = value if label == conclusion_label else structural(row, names[label])
                if not _text(field_value):
                    issues.append({"code": "research_field_missing", "sheet": ws.title,
                                   "cell": ws.cell(row, names[label]).coordinate, "field": label})
                    missing = True
            if missing:
                continue
            hidden = ws.sheet_state != "visible" or bool(ws.row_dimensions[row].hidden)
            table_rows.append(_row(structural(row, names[order_label]), structural(row, names["配置项"]),
                                   _text(value), {"sheet": ws.title, "cell": cell.coordinate},
                                   structural(row, names.get("适用条件与不适用边界", 0)) or "",
                                   structural(row, names.get("建议材料 / 验证依据", 0)) or "", hidden=hidden))
            record = table_rows[-1]
            if '调研子项' in names:
                record['research_id'] = _text(structural(row, names['调研子项']))
            if conclusion_label == '答案':
                record['scope_fields'] = {label: _text(structural(row, names.get(label)))
                                          for label in ('适用公司/院区', '适用业务范围', '上线阶段/期间')}
                record['applicability'] = '；'.join(f'{label}：{value}' for label, value in record['scope_fields'].items() if value)
                record['suggested_evidence'] = _text(structural(row, names.get('依据/来源')))
                record['attachment_materials'] = _text(structural(row, names.get('附件/补充材料')))
        tables.append((conclusion_label, table_rows))
    workbook.close()
    cached.close()
    # 保留每张结果表；镜像或冲突由定位回链核对，不因表头相似而丢弃来源。
    rows = [record for _label, table in tables for record in table]
    if len(tables) > 1:
        issues.append({"code": "multiple_conclusion_tables", "count": len(tables),
                       "selected": "全部结论表", "skipped": 0})
    return {"rows": rows, "worksheets": worksheets, "issues": issues,
            'unanswered': unanswered, 'pending_locations': pending_locations}


def parse_docx(path: Path) -> dict:
    import docx

    document = docx.Document(str(path))
    rows, issues = [], []
    order = item = ""
    for block in iter_block_items(document):
        if block.__class__.__name__ == "Paragraph":
            text = _text(block.text)
            if not text:
                continue
            heading = re.match(r"^第?([一二三四五六七八九十\d]+)[章节、.]\s*(.+)$", text) or \
                      re.match(r"^(\d+(?:\.\d+)+)\s+(.+)$", text)
            if heading:
                order, item = heading.group(1), heading.group(2)[:60]
            elif re.search(r"结论|结果[:：]|答[:：]", text) or len(text) > 15:
                rows.append(_row(order, item, text, {"paragraph": text[:30]}))
        else:  # table
            for r_idx, table_row in enumerate(block.rows, 1):
                cells = [_text(c.text) for c in table_row.cells]
                if len(cells) >= 3 and any(cells):
                    rows.append(_row(cells[0] or order, cells[1] or item, cells[-1],
                                     {"table_row": r_idx}, *(cells[2:-1] + [""] * 2)[:2]))
    return {"rows": rows, "worksheets": [], "issues": issues}


def iter_block_items(parent):
    from docx.document import Document as _Doc
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = parent.element.body if isinstance(parent, _Doc) else parent._element
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def parse_pdf(path: Path) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    rows, issues = [], []
    for number, page in enumerate(reader.pages, 1):
        text = _text(page.extract_text())
        if not text:
            issues.append({"code": "page_without_text_layer", "page": number})
            continue
        order = ""
        for line in text.splitlines():
            if not line.strip():
                continue
            heading = re.match(r"^(\d+(?:[.、．]\d+)*)[、.\s]\s*(.{2,40})$", line.strip())
            if heading and len(line.strip()) < 50:
                order, item = heading.group(1), heading.group(2)
                continue
            if re.search(r"结论|结果[:：]|答[:：]|：", line) or order:
                rows.append(_row(order, item if order else line[:20], line,
                                 {"page": number, "excerpt": line[:30]}))
    return {"rows": rows, "worksheets": [], "issues": issues}


def parse_plain(path: Path) -> dict:
    text = _unicode_text(path)
    lines = text.splitlines()
    # 同时出现“问：”与“答：”结构时按 QA 块解析；否则保持逐行清单模式。
    if any(QUESTION.match(line) for line in lines) and any(ANSWER.match(line) for line in lines):
        return _parse_qa_lines(lines, path.parent)
    rows = []
    order = item = ""
    for number, line in enumerate(lines, 1):
        line = line.strip()
        if not line:
            continue
        if line.startswith("#"):
            heading = re.sub(r"^#+\s*", "", line)
            order, item = ("", heading[:60]) if not re.match(r"^\d", heading) else (re.match(r"^(\S+?)[\s、.]", heading).group(1), heading)
            continue
        if line.startswith(("|", "---")):
            continue
        rows.append(_row(order, item, line, {"line": number}))
    return {"rows": rows, "worksheets": [], "issues": []}


def _parse_qa_lines(lines: list[str], base_dir: Path) -> dict:
    rows, issues = [], []
    state = None  # {"question", "answer", "attachments"}

    def flush():
        nonlocal state
        if state and state["answer"].strip():
            evidence, refs = _attachments(state["attachments"], issues, base_dir)
            rows.append(_row("", state["question"][:80], state["answer"].strip(),
                             {"question": state["question"], "kind": "qa"},
                             "", "；".join(evidence), detail_refs=refs))
        state = None

    for line in lines:
        line = line.rstrip()
        question, answer, attachment = QUESTION.match(line), ANSWER.match(line), ATTACHMENT.match(line)
        if question:
            flush()
            state = {"question": question.group(1).strip(), "answer": "", "attachments": []}
        elif state is None:
            continue
        elif attachment:
            state["attachments"].append(attachment.group(1).strip())
        elif answer:
            state["answer"] += ("\n" if state["answer"] else "") + answer.group(1).strip()
        elif line.strip():
            state["answer"] += ("\n" if state["answer"] else "") + line.strip()
    flush()
    return {"rows": rows, "worksheets": [], "issues": issues}


def _attachments(values: list, issues: list, base_dir: Path) -> tuple[list[str], list[dict]]:
    evidence, refs = [], []
    for value in values:
        if not isinstance(value, str) or not value.strip():
            continue
        evidence_one, ref = _attachment(value.strip(), issues, base_dir)
        if ref is None:
            continue
        evidence.append(evidence_one)
        refs.append(ref)
    return evidence, refs


def _attachment(value: str, issues: list, base_dir: Path) -> tuple[str, dict | None]:
    """附件是链接或本地路径；本工具不做任何网络请求。

    链接只做格式校验（http/https、无凭据）并原样保留为验证依据；需要解析链接内容时，
    由用户把附件下载到本地后在 QA 中引用本地路径，本地文件用既有解析器读取。
    """
    try:
        parsed = urlsplit(value)
        # POSIX 允许 // 开头的本地路径；仅有 authority 不足以把它判成远端链接。
        is_url = parsed.scheme in {"http", "https"} or parsed.username is not None or parsed.password is not None
        rejected = is_url and (parsed.username is not None or parsed.password is not None
                               or parsed.scheme not in {"http", "https"} or not parsed.hostname
                               or any(ord(c) < 32 for c in value))
    except ValueError:
        rejected = True
        is_url = True
    if rejected:
        issues.append({"code": "attachment_rejected", "reason": "链接包含凭据或不是有效链接"})
        return "", None
    if is_url:
        return value, {"attachment": value, "fetched": False,
                       "note": "链接保留为验证依据；解析内容需下载到本地后引用本地路径"}
    path = Path(value)
    if not path.is_absolute():
        path = base_dir / path
    if not path.is_file():
        issues.append({"code": "attachment_missing", "reason": "附件文件不存在"})
        return value, {"attachment": value, "fetched": False}
    parser = PARSERS.get(path.suffix.lower())
    if parser is None:
        issues.append({"code": "attachment_unsupported_type", "suffix": path.suffix})
        return str(path), {"attachment": str(path), "local": True, "fetched": False}
    try:
        parsed = parser(path)
    except Exception as exc:  # 损坏文件会抛各类异常；保留引用并记录问题，不中断整批解析。
        issues.append({"code": "attachment_parse_failed", "reason": type(exc).__name__})
        return str(path), {"attachment": str(path), "local": True, "fetched": False}
    locations = [row.get("locator") for row in parsed.get("rows", []) if row.get("locator")]
    return str(path), {"attachment": str(path), "local": True, "fetched": True, "parser": path.suffix.lower(),
                       "candidate_issues": parsed.get("issues", []),
                       "rows": len(parsed["rows"]), "locations": locations[:50],
                       "location_count": len(locations), "locations_truncated": len(locations) > 50,
                       "sample": [row.get("conclusion", "")[:60] for row in parsed["rows"][:5]]}


def parse_qa_json(path: Path) -> dict:
    """QA JSON：一个问题 + 一个回答（文本 + 附件链接）；问题映射矩阵键由确认后的 mapping 完成。"""
    data = json.loads(_unicode_text(path))
    items = data.get("items") if isinstance(data, dict) else data
    if not isinstance(items, list):
        raise ValueError("QA JSON 需为对象列表或 {\"items\": [...]}")
    rows, issues = [], []
    for index, item in enumerate(items, 1):
        if not isinstance(item, dict):
            issues.append({"code": "qa_item_not_object", "index": index})
            continue
        question, answer = _text(item.get("question")), _text(item.get("answer"))
        if not question or not answer:
            issues.append({"code": "qa_item_missing_question_or_answer", "index": index})
            continue
        attachments = item.get("attachments") or []
        if not isinstance(attachments, list):
            issues.append({"code": "qa_attachments_invalid", "index": index})
            attachments = []
        evidence, refs = _attachments(attachments, issues, path.parent)
        row = _row("", question[:80], answer,
                         {"question": question, "kind": "qa", "index": index},
                         _text(item.get("scope")), "；".join(evidence), detail_refs=refs)
        row["scope"] = _text(item.get("scope"))
        rows.append(row)
    return {"rows": rows, "worksheets": [], "issues": issues}


PARSERS = {".xlsx": parse_xlsx, ".docx": parse_docx, ".pdf": parse_pdf,
           ".md": parse_plain, ".markdown": parse_plain, ".txt": parse_plain,
           ".json": parse_qa_json}


def parse_source(path: Path, module: str) -> dict:
    parser = PARSERS.get(path.suffix.lower())
    if parser is None:
        extracted = {"rows": [], "issues": [{"code": "candidate_reader_unavailable",
                      "suffix": path.suffix.lower(), "blocking": True}]}
    else:
        try:
            extracted = parser(path)
        except Exception as exc:
            extracted = {"rows": [], "issues": [{"code": "candidate_parse_failed",
                          "error_type": type(exc).__name__, "blocking": True}]}
    return {"schema": "dy-research-source/v1", "module": module,
            "source_file": path.name, "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "records": extracted["rows"], "worksheets": extracted.get("worksheets", []),
            "issues": extracted.get("issues", []),
            'unanswered': extracted.get('unanswered', []),
            'pending_locations': extracted.get('pending_locations', []),
            "coverage": {"source_subitems": len(extracted["rows"]),
                         'source_pending_cells': len(extracted.get('pending_locations', [])),
                         "contains_pending_questions": sum(r["contains_pending_questions"] for r in extracted["rows"]),
                         "confirmed_written": 0},
            "limitations": "解析候选事实不是业务确认；不自动写 confirmed、不改验收矩阵、不把来源要求当系统现状。"}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--module", required=True)
    parser.add_argument("--workspace", type=Path, required=True)
    parser.add_argument("--force", action="store_true", help="显式创建新的解析轮次并保留历史")
    args = parser.parse_args()
    if not args.workspace.is_dir():
        parser.error("工作区必须是已存在的目录")
    result = parse_source(args.source, args.module)
    attachments = []
    archived_attachments = {}
    for record in result["records"]:
        for reference in record.get("detail_refs", []):
            attachment = reference.get("attachment")
            if not (reference.get("local") or reference.get("fetched")) or not attachment:
                if attachment:
                    attachments.append({"name": Path(urlsplit(attachment).path).name or "remote-reference",
                                        "status": "not_local", "path": None})
                continue
            source = Path(attachment)
            identity = source.resolve()
            candidate_issues = reference.pop("candidate_issues", [])
            if identity not in archived_attachments:
                archived = archive_source(args.workspace, source, args.module, force=args.force,
                                          candidates={"issues": candidate_issues})
                archived_attachments[identity] = archived
                child_issues = json.loads((args.workspace / archived["path"] / "issues.json").read_text(encoding="utf-8"))
                result["issues"].extend({**issue, "source_component": "attachment_archive",
                                         "attachment": source.name,
                                         "source_sha256": archived["source_sha256"]}
                                        for issue in child_issues if issue.get("blocking", True))
            else:
                archived = archived_attachments[identity]
            reference.update({"attachment": source.name, "archive_path": archived["path"],
                              "source_sha256": archived["source_sha256"], "status": archived["status"]})
            reference.pop("sample", None)
            attachments.append({"name": source.name, "status": archived["status"],
                                "path": archived["path"], "issues": archived["coverage"]["unresolved"]})
    # 同一附件可被多个 QA 引用，但同一阻塞缺口只进入汇总一次。
    unique_issues = {}
    for issue in result["issues"]:
        unique_issues.setdefault(json.dumps(issue, ensure_ascii=False, sort_keys=True), issue)
    result["issues"] = list(unique_issues.values())
    source_archive = archive_source(args.workspace, args.source, args.module, force=args.force,
                                    candidates=result)
    print(json.dumps({"module": args.module, "status": source_archive["status"],
                      "path": source_archive["path"], "source_sha256": source_archive["source_sha256"],
                      "files": {name: f"{source_archive['path']}/{name}.json"
                                for name in ('manifest', 'candidates', 'coverage', 'issues', 'content')},
                      "pending": {"candidate_records": result["coverage"]["contains_pending_questions"],
                                  "source_cells": result["coverage"]["source_pending_cells"],
                                  "unanswered": len(result["unanswered"])},
                      "source_subitems": result["coverage"]["source_subitems"],
                      "issues": source_archive["coverage"]["unresolved"],
                      "worksheets": len(result["worksheets"]), "attachments": attachments},
                     ensure_ascii=False))


if __name__ == "__main__":
    main()
